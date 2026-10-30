"""
Single Python entrypoint for Vercel.

As of mid-2026, Vercel's Python runtime expects one entrypoint app (Flask,
FastAPI, Django, etc.) rather than treating every file in /api as its own
function. This file is that entrypoint — it's a small Flask app with routes
for each feature:

  GET  /api/recent-logs         -> last N days of Daily Job Logs from Dropbox
  GET/POST /api/typeform-webhook -> real-time Typeform submission receiver
  GET  /api/upload-daily-logs   -> legacy/backup: manually re-run the Gmail scan
  POST /api/publish             -> admin publishes approved items to Calvin's dashboard
  GET  /api/published           -> Calvin's dashboard reads the latest published state
  POST /api/add-note            -> notes and flags from either side, appended in place

See vercel.json for the rewrite rule that routes /api/* here.
"""

import os
import sys
import re
import json
import io
import csv
import hmac
import hashlib
import base64
import datetime
import time
from functools import wraps

import requests
import dropbox
from concurrent.futures import ThreadPoolExecutor
import openpyxl
from flask import Flask, request, jsonify, session, send_from_directory

sys.path.append(os.path.join(os.path.dirname(__file__), "..", "scripts"))
from dropbox_pdf_utils import (  # noqa: E402
    get_dropbox_client,
    format_date_for_jacque,
    upload_pdf_and_photos,
    get_recent_daily_logs,
    list_log_index,
    get_file_links,
    extract_text_from_pdf_bytes,
    DAYS_TO_SHOW_DEFAULT,
    DROPBOX_BASE_FOLDER,
    MAX_RANGE_DAYS,
)
from daily_log_uploader import run_gmail_djl_uploader  # noqa: E402
import notion_utils  # noqa: E402
import form_parsers  # noqa: E402
import djl_titles  # noqa: E402

app = Flask(__name__)
# Falls back to a fixed dev key so the app doesn't crash if this isn't set yet —
# but sessions are only truly secure once FLASK_SECRET_KEY is set in Vercel.
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "dev-insecure-key-change-me-in-vercel")
app.permanent_session_lifetime = datetime.timedelta(days=14)

# Repo root — where index.html and dashboard.html live, one level up from /api.
_REPO_ROOT = os.path.join(os.path.dirname(__file__), "..")


# ==========================================
# Static pages — served directly by Flask, since Vercel's current Python
# runtime routes all traffic through this app rather than serving these
# as separate static files.
# ==========================================

@app.route("/")
def root_page():
    return send_from_directory(_REPO_ROOT, "index.html")


@app.route("/index.html")
def admin_page():
    return send_from_directory(_REPO_ROOT, "index.html")


@app.route("/dashboard.html")
def dashboard_page():
    return send_from_directory(_REPO_ROOT, "dashboard.html")


@app.route("/fabshop-portal.html")
def fabshop_portal_page():
    return send_from_directory(_REPO_ROOT, "fabshop-portal.html")


DASHBOARD_STATE_PATH = os.environ.get(
    "DASHBOARD_STATE_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/published.json"
)
FABSHOP_REVENUE_PATH = os.environ.get(
    "FABSHOP_REVENUE_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/fabshop_revenue.json"
)
FABSHOP_DAILY_ENTRIES_PATH = os.environ.get(
    "FABSHOP_DAILY_ENTRIES_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/fabshop_daily_entries.json"
)
FABSHOP_EMPLOYEES_PATH = os.environ.get(
    "FABSHOP_EMPLOYEES_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/fabshop_employees.json"
)
FABSHOP_LEADS_PATH = os.environ.get(
    "FABSHOP_LEADS_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/fabshop_leads.json"
)
DAILY_LOG_NOTES_PATH = os.environ.get(
    "DAILY_LOG_NOTES_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/daily_log_notes.json"
)
DAILY_LOG_TITLE_OVERRIDES_PATH = os.environ.get(
    "DAILY_LOG_TITLE_OVERRIDES_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/daily_log_title_overrides.json"
)
FIELD_ID_ROLE_MAP_PATH = os.environ.get(
    "FIELD_ID_ROLE_MAP_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/typeform_field_role_map.json"
)

NOTION_BATCH_REPORTS_DATASOURCE_ID = os.environ.get(
    "NOTION_BATCH_REPORTS_DATASOURCE_ID", "ec5dda0d-7c82-4c4a-a613-fa28c7702c9c"
)
# Plant Breakdown Material Tracker (McCook Concrete Operations Hub). Filled by the
# MCI driver app: when the batch manager ticks "Plant Breakdown" on the daily
# report, one row per material lost (Material, Amount, Material Type, Report
# Date, Batch Manager, Daily Report Ref).
NOTION_PLANT_BREAKDOWN_DATASOURCE_ID = os.environ.get(
    "NOTION_PLANT_BREAKDOWN_DATASOURCE_ID", "69468d28-40ff-4dcb-add7-5bf97ff90dd2"
)
NOTION_GRAVEL_SALES_DATASOURCE_ID = os.environ.get(
    "NOTION_GRAVEL_SALES_DATASOURCE_ID", "80c0157b-05b2-46f8-9c19-b6f63c6e99af"
)
NOTION_PURCHASE_ORDERS_DATASOURCE_ID = os.environ.get(
    "NOTION_PURCHASE_ORDERS_DATASOURCE_ID", "7d417736-3252-48b3-8335-4f534905085f"
)
# "Current Job Analyses" is a Notion PAGE (not a database) — a running list
# of links to per-job SharePoint budget workbooks, organized by week.
NOTION_JOB_ANALYSES_PAGE_ID = os.environ.get(
    "NOTION_JOB_ANALYSES_PAGE_ID", "1994562c-e56f-8092-af5e-d751067576c1"
)
# "Typeform Leads. FEED HERE" — the sales pipeline / CRM database. This is
# also where Kevin's reports land, since he's one of the Sales Reps.
NOTION_SALES_PIPELINE_DATASOURCE_ID = os.environ.get(
    "NOTION_SALES_PIPELINE_DATASOURCE_ID", "df4a2436-3221-4e35-921f-2c78440188e9"
)
# Kevin's "Daily Management Form" — syncs directly into Notion via a
# separate Typeform-to-Notion connector, unrelated to our own script/webhook.
NOTION_KEVIN_FORM_DATASOURCE_ID = os.environ.get(
    "NOTION_KEVIN_FORM_DATASOURCE_ID", "3ab4562c-e56f-8029-b87a-000b05352beb"
)
# Work Orders lives inside a multi-source database. This is the ID of the
# specific "Service Request" data source within it (see notion_utils.query_data_source).
NOTION_WORKORDERS_DATASOURCE_ID = os.environ.get(
    "NOTION_WORKORDERS_DATASOURCE_ID", "52e9dbff-b6db-4374-a77b-31b86c8ce5eb"
)
# Current Jobs is an inline database within a page (not a standalone
# top-level database), so it uses the newer data-sources endpoint like
# Work Orders does.
NOTION_CURRENT_JOBS_DATASOURCE_ID = os.environ.get(
    "NOTION_CURRENT_JOBS_DATASOURCE_ID", "92caa4f5-2f8f-4d5b-9c6f-c1b937685d4f"
)


# ==========================================
# AUTH — two accounts: admin (you) and calvin
# ==========================================

def require_role(*allowed_roles):
    def decorator(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            role = session.get("role")
            if role not in allowed_roles:
                return jsonify({"status": "error", "message": "Not logged in or not authorized"}), 401
            return fn(*args, **kwargs)
        return wrapper
    return decorator


def require_role_or_cron(*allowed_roles):
    """Same as require_role, but also lets Vercel's scheduled Cron trigger
    through — Vercel sends 'Authorization: Bearer <CRON_SECRET>' on cron
    invocations, which carry no login session at all."""
    def decorator(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            cron_secret = os.environ.get("CRON_SECRET")
            auth_header = request.headers.get("Authorization", "")
            if cron_secret and auth_header == f"Bearer {cron_secret}":
                return fn(*args, **kwargs)
            role = session.get("role")
            if role not in allowed_roles:
                return jsonify({"status": "error", "message": "Not logged in or not authorized"}), 401
            return fn(*args, **kwargs)
        return wrapper
    return decorator


@app.route("/api/login", methods=["POST"])
def login():
    """
    No credentials — just a role picker. Body: { role: "admin" | "calvin" }.
    This is intentionally not real authentication: anyone with the link can
    pick either role. Fine for now since these URLs aren't shared publicly,
    but worth adding real credentials later if that changes.
    """
    body = request.get_json(force=True) or {}
    role = body.get("role")

    if role not in ("admin", "calvin"):
        return jsonify({"status": "error", "message": "Invalid role"}), 400

    session.permanent = True
    session["role"] = role
    return jsonify({"status": "ok", "role": role})


@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"status": "ok"})


@app.route("/api/session", methods=["GET"])
def get_session():
    return jsonify({"logged_in": "role" in session, "role": session.get("role")})


# ==========================================
# /api/recent-logs
# ==========================================

def _daily_log_key(log):
    """Stable key for a Dropbox-derived log entry, since these don't have a
    real ID of their own (they're just files in a folder each time)."""
    return f"{log.get('date', '')}|{log.get('name', '')}|{log.get('job_or_wo') or ''}"


def _notes_for_log(all_notes, log):
    """Exact key match first; falls back to any note sharing just date+name,
    since a bulk-imported note's job/WO tag won't always exactly match
    whatever tag ended up in the real Dropbox filename."""
    exact = all_notes.get(_daily_log_key(log))
    if exact:
        return exact
    prefix = f"{log.get('date', '')}|{log.get('name', '')}|"
    combined = []
    for key, notes in all_notes.items():
        if key.startswith(prefix):
            combined.extend(notes)
    return combined


@app.route("/api/recent-logs", methods=["GET"])
@require_role("admin", "calvin")
def recent_logs():
    try:
        days = int(request.args.get("days", DAYS_TO_SHOW_DEFAULT))
        # Optional date filter (Calvin's dashboard): ?start=YYYY-MM-DD&end=YYYY-MM-DD
        def _iso(name):
            v = (request.args.get(name) or "").strip()
            try:
                return datetime.date.fromisoformat(v) if v else None
            except ValueError:
                raise ValueError(f"'{name}' must be a date like 2026-10-01")
        start, end = _iso("start"), _iso("end")
        logs = get_recent_daily_logs(days=days, start=start, end=end)
        _apply_log_corrections(get_dropbox_client(), logs)
        return jsonify({"days": days, "count": len(logs), "logs": logs})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


def _apply_log_corrections(dbx, logs):
    """Admin title corrections and manual notes, applied to logs from Dropbox."""
    all_notes = _read_json_from_dropbox(dbx, DAILY_LOG_NOTES_PATH, {})
    title_overrides = _read_json_from_dropbox(dbx, DAILY_LOG_TITLE_OVERRIDES_PATH, {})

    for log in logs:
        pdf_path = log.get("dropbox_pdf_path")
        if pdf_path and pdf_path in title_overrides:
            override = title_overrides[pdf_path]
            if override.get("name"):
                log["name"] = override["name"]
            if override.get("job_or_wo") is not None:
                log["job_or_wo"] = override["job_or_wo"]
            log["title_corrected"] = True
        log["manual_notes"] = _notes_for_log(all_notes, log)
    return logs


@app.route("/api/log-links", methods=["POST"])
@require_role("admin", "calvin")
def log_links():
    """Links for log files, made when someone clicks (not when the page
    loads). Body: {"paths": [...], "mode": "view" | "download" | "temp"}.
    The browser then fetches the file straight from Dropbox - nothing passes
    through this server."""
    try:
        body = request.get_json(force=True) or {}
        paths = body.get("paths") or ([body["path"]] if body.get("path") else [])
        mode = body.get("mode", "view")
        if mode not in ("view", "download", "temp"):
            return jsonify({"status": "error", "message": "Bad mode"}), 400
        if not paths or len(paths) > 30:
            return jsonify({"status": "error", "message": "Between 1 and 30 paths required"}), 400
        links = get_file_links(get_dropbox_client(), paths, mode)
        return jsonify({"status": "ok", "links": links})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/daily-log-notes", methods=["POST"])
@require_role("admin", "calvin")
def add_daily_log_note():
    """A manual comment/correction on a Dropbox-derived log entry — a
    stopgap for entries whose real text wasn't captured at upload time."""
    try:
        body = request.get_json(force=True)
        date = body.get("date", "")
        name = body.get("name", "")
        job_or_wo = body.get("job_or_wo") or ""
        text = (body.get("text") or "").strip()
        if not text:
            return jsonify({"status": "error", "message": "Note text is required"}), 400

        key = f"{date}|{name}|{job_or_wo}"
        dbx = get_dropbox_client()
        all_notes = _read_json_from_dropbox(dbx, DAILY_LOG_NOTES_PATH, {})
        all_notes.setdefault(key, []).append({
            "author": body.get("author") or session.get("role", "admin"),
            "text": text,
            "created_at": datetime.datetime.utcnow().isoformat() + "Z",
        })
        _write_json_to_dropbox(dbx, DAILY_LOG_NOTES_PATH, all_notes)
        return jsonify({"status": "ok", "notes": all_notes[key]})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


def _parse_submit_timestamp(val):
    """Typeform CSV exports 'Submit Date (UTC)' as 'YYYY-MM-DD HH:MM:SS'."""
    if not val:
        return None
    try:
        return datetime.datetime.strptime(val.strip(), "%Y-%m-%d %H:%M:%S").isoformat()
    except ValueError:
        return None


def _parse_daily_job_log_csv_row(row):
    """
    Parses one row of a Typeform CSV export of the real 'Daily Job Log' form
    (the one with an actual Job Number field — not the management-log or
    Fab Shop forms). Column names come directly from the export, so this is
    more reliable than the webhook's keyword-guessing on field titles.
    """
    name = (row.get("Your Name") or "").strip()
    raw_date = (row.get("*Date*") or "")[:10]
    try:
        date_iso = datetime.date.fromisoformat(raw_date).isoformat() if raw_date else None
    except ValueError:
        date_iso = None

    job_field = (row.get("*Job number* ") or row.get("*Job number*") or "").strip()
    wo_num = (row.get("Workorder #") or "").strip()
    wo_name = (row.get("Workorder Name") or "").strip()

    job_tags = _extract_job_numbers(job_field)
    if not job_tags:
        job_tags = _extract_job_numbers(wo_num)
    job_wo = ", ".join(job_tags) if job_tags else _clean_or_none(job_field)
    if not job_wo and wo_name:
        job_wo = wo_name

    lines = []
    work_desc = _clean_or_none(row.get("What did your team work on today?"))
    if work_desc:
        lines.append(f"Work: {work_desc}")
    productivity = _clean_or_none(row.get("*Productivity Issues*"))
    if productivity:
        lines.append(f"Productivity issues: {productivity}")
    incident = _clean_or_none(row.get("Incident Report?")) or _clean_or_none(row.get("*Safety Incident*"))
    if incident:
        lines.append(f"Incident: {incident}")
    todo = _clean_or_none(row.get("To do:"))
    if todo:
        lines.append(f"To do: {todo}")
    scope_desc = _clean_or_none(row.get("Please describe additional work:"))
    if scope_desc:
        lines.append(f"Additional work: {scope_desc}")

    return {
        "name": name,
        "date": date_iso,
        "job_or_wo": job_wo,
        "text": "\n".join(lines),
        "submitted_at": _parse_submit_timestamp(row.get("Submit Date (UTC)")),
    }


def _parse_management_log_csv(content):
    """
    Parses the Management Log CSV export positionally rather than by column
    name — several columns (e.g. every 'How long did that take?') repeat
    with the exact same header text, which would silently lose data under
    a name-keyed dict. Also matches the real typo 'Next Projet:' on one slot.
    """
    reader = csv.reader(io.StringIO(content))
    header = next(reader)

    name_idx = None
    date_idx = None
    submit_idx = None
    slot_cols = []  # (index, kind) for every classified project-slot column

    for i, h in enumerate(header):
        hl = h.lower().strip()
        if hl.startswith("your name"):
            name_idx = i
        elif "submit date" in hl:
            submit_idx = i
        elif "date" in hl and name_idx is not None and date_idx is None and i <= name_idx + 2:
            date_idx = i
        else:
            kind = _classify_field(h)
            if kind:
                slot_cols.append((i, kind))

    entries = []
    for row in reader:
        if not row or (name_idx is not None and name_idx >= len(row)):
            continue
        name = (row[name_idx] or "").strip() if name_idx is not None else ""
        raw_date = (row[date_idx] or "")[:10] if date_idx is not None and date_idx < len(row) else ""
        try:
            date_iso = datetime.date.fromisoformat(raw_date).isoformat() if raw_date else None
        except ValueError:
            date_iso = None
        submitted_at = _parse_submit_timestamp(row[submit_idx]) if submit_idx is not None and submit_idx < len(row) else None
        if not name or not date_iso:
            continue

        projects = []
        current = None
        for idx, kind in slot_cols:
            if idx >= len(row):
                continue
            val = (row[idx] or "").strip()
            if kind == "project":
                if not val:
                    current = None
                    continue
                current = {"project": val, "did": "", "duration": "", "next_action": ""}
                projects.append(current)
            elif kind == "did" and current is not None:
                current["did"] = val
            elif kind == "duration" and current is not None:
                current["duration"] = val
            elif kind == "next_action" and current is not None:
                current["next_action"] = val
            # "continuation" (the yes/no "any more projects?") is skipped

        job_wo = None
        for p in projects:
            job_wo = _extract_job_number(p["project"])
            if job_wo:
                break
        if not job_wo:
            job_wo = "Daily Management Log"

        lines = []
        for p in projects:
            if not (p["project"] or p["did"]):
                continue
            line = f"{p['project'] or 'Project'}: {p['did']}"
            if p["duration"]:
                line += f" ({p['duration']})"
            next_clean = _clean_or_none(p["next_action"])
            if next_clean:
                line += f" — Next: {next_clean}"
            lines.append(line)

        entries.append({
            "name": name,
            "date": date_iso,
            "job_or_wo": job_wo,
            "text": "\n".join(lines),
            "submitted_at": submitted_at,
        })

    return entries


def _detect_csv_type(header_row):
    joined = " | ".join(h.lower().replace("*", "") for h in header_row)
    if "job number" in joined:
        return "daily_job_log"
    if "how long did that take" in joined and "next action" in joined:
        return "management_log"
    return "unknown"


@app.route("/api/bulk-import-daily-logs", methods=["POST"])
@require_role("admin")
def bulk_import_daily_logs():
    """
    Upload a Typeform CSV export (Daily Job Log or Management Log — the
    type is auto-detected from the header row) to backfill description
    text on existing entries. Stored the same way as a manual note, tagged
    'CSV Import'. Safe to re-run: skips rows whose exact text was already
    imported.
    """
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "No file included in upload"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"status": "error", "message": "No file selected"}), 400

    try:
        content = file.read().decode("utf-8-sig")

        # Peek the header to figure out which form this export is from
        header_row = next(csv.reader(io.StringIO(content)))
        csv_type = _detect_csv_type(header_row)

        if csv_type == "daily_job_log":
            reader = csv.DictReader(io.StringIO(content))
            parsed_rows = [_parse_daily_job_log_csv_row(row) for row in reader]
        elif csv_type == "management_log":
            parsed_rows = _parse_management_log_csv(content)
        else:
            return jsonify({
                "status": "error",
                "message": "Couldn't recognize this CSV's form type (no Job Number field, and no repeating project-slot pattern found). Let Claude know what form this is from so support for it can be added."
            }), 400

        dbx = get_dropbox_client()
        all_notes = _read_json_from_dropbox(dbx, DAILY_LOG_NOTES_PATH, {})

        imported = 0
        skipped_no_data = 0
        skipped_duplicate = 0

        for parsed in parsed_rows:
            if not parsed["name"] or not parsed["date"] or not parsed["text"]:
                skipped_no_data += 1
                continue

            key = f"{parsed['date']}|{parsed['name']}|{parsed['job_or_wo'] or ''}"
            existing = all_notes.get(key, [])
            if any(n.get("text") == parsed["text"] for n in existing):
                skipped_duplicate += 1
                continue

            existing.append({
                "author": "CSV Import",
                "text": parsed["text"],
                "created_at": datetime.datetime.utcnow().isoformat() + "Z",
            })
            all_notes[key] = existing
            imported += 1

        _write_json_to_dropbox(dbx, DAILY_LOG_NOTES_PATH, all_notes)
        return jsonify({
            "status": "ok",
            "csv_type": csv_type,
            "imported": imported,
            "skipped_no_data": skipped_no_data,
            "skipped_duplicate": skipped_duplicate,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/bulk-fix-titles", methods=["POST"])
@require_role("admin")
def bulk_fix_titles():
    """
    Corrects the TITLE (name + job/WO) shown on existing entries whose name
    came through as 'Unknown' — unlike the description-text import, this
    changes what's actually displayed at the top of the card. Matching is
    done by closest submit-timestamp on the same date, since a broken
    entry's own name/job fields can't be used to find its real CSV row.
    Best-effort: please review the results rather than trusting blindly.
    """
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "No file included in upload"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"status": "error", "message": "No file selected"}), 400

    try:
        content = file.read().decode("utf-8-sig")
        header_row = next(csv.reader(io.StringIO(content)))
        csv_type = _detect_csv_type(header_row)

        if csv_type == "daily_job_log":
            reader = csv.DictReader(io.StringIO(content))
            parsed_rows = [_parse_daily_job_log_csv_row(row) for row in reader]
        elif csv_type == "management_log":
            parsed_rows = _parse_management_log_csv(content)
        else:
            return jsonify({
                "status": "error",
                "message": "Couldn't recognize this CSV's form type."
            }), 400

        parsed_rows = [r for r in parsed_rows if r.get("submitted_at") and r.get("date")]

        # Wide window so this can catch older broken entries too, not just
        # whatever the dashboard's normal 5-day view shows.
        logs = get_recent_daily_logs(days=60)
        broken = [
            l for l in logs
            if (l.get("name") or "").strip().lower() == "unknown"
            and l.get("uploaded_at") and l.get("dropbox_pdf_path")
        ]

        dbx = get_dropbox_client()
        overrides = _read_json_from_dropbox(dbx, DAILY_LOG_TITLE_OVERRIDES_PATH, {})

        from collections import defaultdict
        broken_by_date = defaultdict(list)
        for b in broken:
            broken_by_date[b["date"]].append(b)
        csv_by_date = defaultdict(list)
        for i, r in enumerate(parsed_rows):
            csv_by_date[r["date"]].append((i, r))

        matched = 0
        unmatched = 0
        used_csv_rows = set()
        match_details = []

        for date, broken_entries in broken_by_date.items():
            candidates = csv_by_date.get(date, [])
            for b in broken_entries:
                try:
                    b_time = datetime.datetime.fromisoformat(b["uploaded_at"].replace("Z", ""))
                except ValueError:
                    unmatched += 1
                    continue

                best, best_diff = None, None
                for i, r in candidates:
                    if i in used_csv_rows:
                        continue
                    try:
                        r_time = datetime.datetime.fromisoformat(r["submitted_at"])
                    except ValueError:
                        continue
                    diff = abs((b_time - r_time).total_seconds())
                    if best_diff is None or diff < best_diff:
                        best_diff, best = diff, (i, r)

                if best:
                    i, r = best
                    used_csv_rows.add(i)
                    overrides[b["dropbox_pdf_path"]] = {"name": r["name"], "job_or_wo": r["job_or_wo"]}
                    match_details.append({
                        "date": date, "corrected_name": r["name"], "corrected_job_or_wo": r["job_or_wo"],
                        "time_diff_seconds": round(best_diff),
                    })
                    matched += 1
                else:
                    unmatched += 1

        _write_json_to_dropbox(dbx, DAILY_LOG_TITLE_OVERRIDES_PATH, overrides)
        return jsonify({
            "status": "ok",
            "csv_type": csv_type,
            "total_broken_found": len(broken),
            "matched": matched,
            "unmatched": unmatched,
            "match_details": match_details,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# /api/typeform-webhook
# ==========================================

def verify_signature(raw_body, signature_header):
    secret = os.environ.get("TYPEFORM_WEBHOOK_SECRET")
    if not secret:
        return True  # no secret configured — should be set in production
    if not signature_header or not signature_header.startswith("sha256="):
        return False
    expected = base64.b64encode(
        hmac.new(secret.encode(), raw_body, hashlib.sha256).digest()
    ).decode()
    provided = signature_header.split("sha256=", 1)[1]
    return hmac.compare_digest(expected, provided)


def find_answer(answers, keywords):
    return form_parsers.find_answer(answers, keywords)


def find_answer_with_learning(answers, keywords, role_key, learned_map):
    return form_parsers.find_answer_with_learning(answers, keywords, role_key, learned_map)


def answer_text(ans):
    return form_parsers.answer_text(ans)


def download_file(url):
    token = os.environ.get("TYPEFORM_API_TOKEN")
    headers = {"Authorization": f"Bearer {token}"} if token else {}
    resp = requests.get(url, headers=headers, timeout=20)
    return resp.content if resp.status_code == 200 else None


# ==========================================
# The actual structured parsing logic lives in scripts/form_parsers.py now,
# shared with the Gmail-scan path (scripts/daily_log_uploader.py) so both
# produce identically good output instead of the email path being cruder.
# ==========================================

def _extract_job_numbers(text):
    return form_parsers._extract_job_numbers(text)


def _extract_job_number(text):
    return form_parsers._extract_job_number(text)


def _clean_or_none(text):
    return form_parsers._clean_or_none(text)


def _detect_form_type(answers, learned_map=None):
    return form_parsers.detect_form_type(answers, learned_map)


def _collect_photos(answers):
    photo_bytes_list = []
    for ans in answers:
        if ans.get("type") == "file_url":
            file_bytes = download_file(ans.get("file_url"))
            if file_bytes:
                photo_bytes_list.append(file_bytes)
    return photo_bytes_list


def _parse_management_log(answers, log_name, log_date, learned_map=None):
    return form_parsers.parse_management_log(answers, log_name, log_date, learned_map)


def _parse_daily_job_log(answers, log_name, log_date, learned_map=None):
    return form_parsers.parse_daily_job_log(answers, log_name, log_date, learned_map)


def _parse_fab_shop_log(answers, log_name, log_date, learned_map=None):
    return form_parsers.parse_fab_shop_log(answers, log_name, log_date, learned_map)



def process_submission(payload):
    form_response = payload.get("form_response", {})
    answers = form_response.get("answers", [])
    submitted_at = form_response.get("submitted_at", "")

    if not answers:
        raise ValueError("Submission had no answers")

    # Typeform's webhook payload does NOT embed a question's title inside
    # answers[].field — that only ever has {id, type, ref}. The real
    # titles live in a completely separate array, form_response.definition
    # .fields[], keyed by that same field id. Every previous attempt to
    # read a title straight off an answer was structurally looking in the
    # wrong place — this was never a "title sometimes missing" problem.
    # Build the id -> title lookup once and enrich every answer with its
    # real title before any classification logic runs.
    definition_fields = form_response.get("definition", {}).get("fields", [])
    field_id_to_title = {f.get("id"): f.get("title") for f in definition_fields if f.get("id")}
    for ans in answers:
        field_id = ans.get("field", {}).get("id")
        if field_id and field_id_to_title.get(field_id):
            ans["field"]["title"] = field_id_to_title[field_id]

    dbx = get_dropbox_client()
    learned_map = _read_json_from_dropbox(dbx, FIELD_ID_ROLE_MAP_PATH, {})

    name_ans = find_answer_with_learning(answers, ["your name", "name"], "submitter_name", learned_map)
    date_ans = find_answer_with_learning(answers, ["date"], "submitted_date", learned_map)
    log_name = _clean_or_none(answer_text(name_ans)) or "Unknown"
    raw_date = answer_text(date_ans) or submitted_at[:10] or ""
    log_date = format_date_for_jacque(raw_date)

    form_type = _detect_form_type(answers, learned_map)
    if form_type == "management_log":
        job_wo, pdf_text, extra = _parse_management_log(answers, log_name, log_date, learned_map)
    elif form_type == "daily_job_log":
        job_wo, pdf_text, extra = _parse_daily_job_log(answers, log_name, log_date, learned_map)
    elif form_type == "fab_shop_log":
        job_wo, pdf_text, extra = _parse_fab_shop_log(answers, log_name, log_date, learned_map)
    else:
        form_title = form_response.get("definition", {}).get("title", "")
        if not form_parsers.is_log_submission(form_type, form_title):
            # Not a log form (e.g. a job application or screening form
            # pointed at this webhook). Those hold applicants' personal
            # details, so nothing is written to the Daily Job Logs folders.
            return {"skipped": True, "reason": "not a job log form", "form_type": form_type,
                    "form_title": form_title}
        # A log form whose answers didn't match a known shape - keep it
        # (generic summary) rather than lose the submission.
        job_wo = "Daily Management Log" if "management" in form_title.lower() else ""
        lines = []
        for ans in answers:
            title = ans.get("field", {}).get("title") or "Field (untitled)"
            val = answer_text(ans)
            if val:
                lines.append(f"{title}: {val}")
        pdf_text = "\n".join(lines)
        extra = {"form_type": "unrecognized"}

    _write_json_to_dropbox(dbx, FIELD_ID_ROLE_MAP_PATH, learned_map)

    # The PDF gets the complete log (every question and answer); pdf_text
    # (the short summary) only goes to the .txt sidecar for the dashboard.
    full_text = form_parsers.render_full_log(answers)

    photo_bytes_list = _collect_photos(answers)

    pdf_title_parts = [log_date, log_name]
    if job_wo:
        pdf_title_parts.append(job_wo)
    pdf_title = " | ".join(pdf_title_parts)

    summary = upload_pdf_and_photos(
        dbx, log_date, log_name, job_wo, pdf_text, pdf_title, photo_bytes_list,
        full_text=full_text,
    )
    summary["log_name"] = log_name
    summary["job_wo"] = job_wo
    summary["log_date"] = log_date
    summary["form_type"] = form_type
    summary.update(extra)
    return summary


@app.route("/api/typeform-webhook", methods=["GET", "POST"])
def typeform_webhook():
    if request.method == "GET":
        return jsonify({"status": "ok", "message": "Typeform webhook endpoint is live. POST only."})

    raw_body = request.get_data()
    signature = request.headers.get("Typeform-Signature")
    if not verify_signature(raw_body, signature):
        return jsonify({"status": "error", "message": "Invalid signature"}), 401

    # Daily Job Logs are now written ONLY by the office PC script
    # (local/daily_job_logs.py, from the Typeform notification emails).
    # This endpoint still answers 200 so Typeform doesn't flag the webhook
    # as failing, but it no longer writes anything to Dropbox. Set
    # DJL_WEBHOOK_WRITES=1 in Vercel to switch the old behaviour back on.
    if os.environ.get("DJL_WEBHOOK_WRITES") != "1":
        return jsonify({"status": "ok", "message": "Received; Daily Job Logs are saved by the office PC script."})

    try:
        payload = json.loads(raw_body.decode("utf-8"))
        summary = process_submission(payload)
        return jsonify({"status": "ok", "summary": summary})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# /api/upload-daily-logs (legacy/backup manual trigger)
# ==========================================

@app.route("/api/upload-daily-logs", methods=["GET", "POST"])
@require_role_or_cron("admin")
def upload_daily_logs():
    # Turned off: Daily Job Logs run on the office PC now (local/daily_job_logs.py),
    # which has no time limit. Set DJL_DASHBOARD_UPLOADS=1 in Vercel to switch it back on.
    if os.environ.get("DJL_DASHBOARD_UPLOADS") != "1":
        return jsonify({"status": "error",
                        "message": "Daily Job Logs now run on the office PC; the dashboard no longer uploads them."}), 410
    try:
        if request.method == "POST":
            # Admin button — processes exactly one batch; the frontend
            # loops itself, showing live progress between calls.
            body = request.get_json(silent=True) or {}
            start_index = int(body.get("start_index", 0))
            # Optional one-time repair: rebuild=True overwrites existing
            # PDFs/.txt files with the full log; days widens the Gmail
            # search window for that run (capped at 120 days).
            rebuild = bool(body.get("rebuild", False))
            days = body.get("days")
            days = max(1, min(int(days), 120)) if days else None
            result = run_gmail_djl_uploader(start_index=start_index, batch_size=10,
                                            rebuild=rebuild, days=days)
            return jsonify({"status": "ok", **result})
        else:
            # GET (the hourly Cron trigger, or a manual browser visit) —
            # can't be interactively driven the way the button's POST loop
            # is, so this loops through batches itself, server-side,
            # stopping once done or once it's used up a safe chunk of the
            # function's time budget (whatever's left gets caught on the
            # next hourly run).
            start_time = time.time()
            start_index = 0
            combined_summary = {"scanned": 0, "pdfs_uploaded": 0, "photos_uploaded": 0, "skipped": 0}
            total = 0
            done = True
            while True:
                # Each call stops on its own when its time runs out, so
                # hand it only what's left of this request's ~45s.
                remaining = 45 - (time.time() - start_time)
                if remaining < 10:
                    done = False
                    break
                result = run_gmail_djl_uploader(start_index=start_index, batch_size=10,
                                                time_budget=remaining)
                for k in combined_summary:
                    combined_summary[k] += result["summary"][k]
                total = result["total"]
                done = result["done"]
                if done:
                    break
                start_index = result["next_index"]
                if time.time() - start_time > 45:  # safety margin under the 60s limit
                    break
            return jsonify({"status": "ok", "summary": combined_summary, "done": done, "total": total})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Dashboard storage (Dropbox-backed — no separate database needed)
# ==========================================

def _read_json_from_dropbox(dbx, path, default):
    try:
        _, res = dbx.files_download(path)
        return json.loads(res.content)
    except dropbox.exceptions.ApiError:
        return default


def _write_json_to_dropbox(dbx, path, data):
    dbx.files_upload(
        json.dumps(data, indent=2).encode("utf-8"),
        path,
        mode=dropbox.files.WriteMode.overwrite,
    )


def _read_published_state(dbx):
    try:
        _, res = dbx.files_download(DASHBOARD_STATE_PATH)
        return json.loads(res.content)
    except dropbox.exceptions.ApiError:
        return {"published_at": None, "items": []}


def _write_published_state(dbx, state):
    dbx.files_upload(
        json.dumps(state, indent=2).encode("utf-8"),
        DASHBOARD_STATE_PATH,
        mode=dropbox.files.WriteMode.overwrite,
    )


@app.route("/api/publish", methods=["POST"])
@require_role("admin")
def publish():
    try:
        body = request.get_json(force=True)
        new_items = body.get("items", [])

        dbx = get_dropbox_client()
        state = _read_published_state(dbx)

        existing_by_id = {i["id"]: i for i in state.get("items", [])}
        for item in new_items:
            item.setdefault("notes", [])
            existing_by_id[item["id"]] = item  # replace or add

        state = {
            "published_at": datetime.datetime.utcnow().isoformat() + "Z",
            "items": list(existing_by_id.values()),
        }
        _write_published_state(dbx, state)
        return jsonify({"status": "ok", "published_at": state["published_at"], "count": len(new_items)})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/published", methods=["GET"])
@require_role("admin", "calvin")
def published():
    try:
        dbx = get_dropbox_client()
        state = _read_published_state(dbx)
        return jsonify(state)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/add-note", methods=["POST"])
@require_role("admin", "calvin")
def add_note():
    """
    Body: { item_id, text, type ("note" or "flag"), category (optional, for flags) }
    Author is taken from the logged-in session, not the request body, so a
    note can't be spoofed as coming from the other person.
    Appends to the item's notes list and re-saves. Used by both Calvin's
    dashboard (notes + flags) and the admin page (replies to Calvin's notes).
    """
    try:
        body = request.get_json(force=True)
        item_id = body.get("item_id")
        if not item_id:
            return jsonify({"status": "error", "message": "item_id is required"}), 400

        dbx = get_dropbox_client()
        state = _read_published_state(dbx)

        target = next((i for i in state.get("items", []) if i["id"] == item_id), None)
        if not target:
            return jsonify({"status": "error", "message": "item not found in published state"}), 404

        target.setdefault("notes", []).append({
            "author": session.get("role"),
            "text": body.get("text", ""),
            "type": body.get("type", "note"),
            "category": body.get("category"),
            "created_at": datetime.datetime.utcnow().isoformat() + "Z",
        })

        _write_published_state(dbx, state)
        return jsonify({"status": "ok", "item": target})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Work Orders — live feed for the admin review queue. Escalated ones get
# tagged distinctly so the frontend can give them a red border.
# ==========================================

@app.route("/api/notion/work-orders", methods=["GET"])
@require_role("admin")
def notion_work_orders():
    try:
        pages = notion_utils.query_data_source(NOTION_WORKORDERS_DATASOURCE_ID, page_size=100)
        items = []
        for page in pages:
            props = page.get("properties", {})
            wo_num = notion_utils.prop_number(props, "WO #")
            job_num = notion_utils.prop_number(props, "Job #")
            contact = notion_utils.prop_text(props, "Reporting Contact")
            status = notion_utils.prop_select(props, "Work Status")
            priority = notion_utils.prop_select(props, "Service Priority")
            date = notion_utils.prop_date(props, "Date")
            desc = notion_utils.prop_text(props, "Description of Work")
            complaint = notion_utils.prop_text(props, "Customer Complaint")

            is_escalated = (status == "Escalated")
            job_key = f"J{int(job_num)}" if job_num else None

            items.append({
                "id": "notion-wo-" + page["id"],
                "notionUrl": page.get("url"),
                "type": "general",
                "source": "notion-workorders",
                "sourceLabel": "Notion · Work Orders",
                "job": job_key,
                "title": (f"WO {int(wo_num)}" if wo_num else "Work order") + (f" — {contact}" if contact else ""),
                "subtitle": f"{status or 'Unknown'}" + (f" · {priority}" if priority else "") + (f" · {date}" if date else ""),
                "tagClass": "escalated" if is_escalated else ("warn" if priority == "Urgent" else "ok"),
                "tagText": status or "Unknown",
                "summary": desc or complaint or "",
                "status": "pending",
            })
        return jsonify({"count": len(items), "items": items})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Notion sources — MCI Batch Reports & Gravel Sales
# ==========================================

@app.route("/api/notion/batch-reports", methods=["GET"])
@require_role("admin", "calvin")
def notion_batch_reports():
    try:
        pages = notion_utils.query_data_source(NOTION_BATCH_REPORTS_DATASOURCE_ID)
        items = []
        for page in pages:
            props = page.get("properties", {})
            name = notion_utils.prop_text(props, "Name") or "Batch report"
            report_date = notion_utils.prop_date(props, "Report Date") or notion_utils.prop_date(props, "Date")
            issues_raw = notion_utils.prop_text(props, "Issues Presented")
            yards = notion_utils.prop_number(props, "Total Yards Out")
            trips = notion_utils.prop_number(props, "Trips Out")
            submitted_by = notion_utils.prop_text(props, "Submitted By")

            # "N/A", "None", "-", etc. are placeholder values meaning "no issue",
            # not real issue descriptions — treat them as empty.
            issues_clean = (issues_raw or "").strip()
            no_issue_placeholders = {"n/a", "na", "none", "no", "no issues", "-", ""}
            has_issue = issues_clean.lower() not in no_issue_placeholders

            items.append({
                "id": "notion-batch-" + page["id"],
                "notionUrl": page.get("url"),
                "type": "general",
                "source": "notion-batch-reports",
                "sourceLabel": "Notion · Daily Batch Reports",
                "job": None,
                "title": name,
                "subtitle": f"{report_date or 'No date'} · {yards or 0} yds · {trips or 0} trips" + (f" · by {submitted_by}" if submitted_by else ""),
                "tagClass": "warn" if has_issue else "ok",
                "tagText": "Issue reported" if has_issue else "On track",
                "summary": issues_clean if has_issue else "",
                "status": "pending",
            })
        return jsonify({"count": len(items), "items": items})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/notion/plant-breakdowns", methods=["GET"])
@require_role("admin", "calvin")
def notion_plant_breakdowns():
    """
    Material lost to plant breakdowns, for McCook Concrete's "Plant Breakdowns" tab.
    ?start=YYYY-MM-DD&end=YYYY-MM-DD (either may be left out = no limit that side).
    Returns totals per material for the period, plus one entry per breakdown day
    with each material's amount and that day's batch report issues for context.
    """
    try:
        def _iso(name):
            v = (request.args.get(name) or "").strip()
            try:
                return datetime.date.fromisoformat(v).isoformat() if v else None
            except ValueError:
                raise ValueError(f"'{name}' must be a date like 2026-10-01")
        start, end = _iso("start"), _iso("end")

        rows = notion_utils.query_data_source_all(NOTION_PLANT_BREAKDOWN_DATASOURCE_ID)

        days = {}
        for page in rows:
            props = page.get("properties", {})
            day = (notion_utils.prop_date(props, "Report Date") or page.get("created_time", ""))[:10]
            if not day or (start and day < start) or (end and day > end):
                continue
            amount = notion_utils.prop_number(props, "Amount") or 0
            material = notion_utils.prop_text(props, "Material") or "Unknown material"
            entry = days.setdefault(day, {"date": day, "batch_managers": [], "materials": [],
                                          "issues": "", "batch_report_url": None, "notion_urls": []})
            manager = notion_utils.prop_text(props, "Batch Manager")
            if manager and manager not in entry["batch_managers"]:
                entry["batch_managers"].append(manager)
            entry["materials"].append({
                "material": material,
                "amount": amount,
                "type": notion_utils.prop_select(props, "Material Type") or "Standard",
            })
            if page.get("url"):
                entry["notion_urls"].append(page["url"])

        # That day's batch report "Issues Presented" usually says what broke down.
        if days:
            try:
                no_issue = {"n/a", "na", "none", "no", "no issues", "-", ""}
                for page in notion_utils.query_data_source_all(NOTION_BATCH_REPORTS_DATASOURCE_ID):
                    props = page.get("properties", {})
                    day = (notion_utils.prop_date(props, "Report Date") or notion_utils.prop_date(props, "Date") or "")[:10]
                    if day in days:
                        issues = (notion_utils.prop_text(props, "Issues Presented") or "").strip()
                        if issues.lower() not in no_issue and issues not in days[day]["issues"]:
                            days[day]["issues"] = (days[day]["issues"] + "\n" + issues).strip()
                        days[day]["batch_report_url"] = days[day]["batch_report_url"] or page.get("url")
            except Exception:
                pass  # issues are context only; the material numbers still show

        totals = {}
        for entry in days.values():
            merged = {}
            for m in entry["materials"]:
                key = m["material"].strip().lower()
                if key in merged:
                    merged[key]["amount"] += m["amount"]
                else:
                    merged[key] = dict(m)
            entry["materials"] = sorted(merged.values(), key=lambda m: -m["amount"])
            entry["total_amount"] = round(sum(m["amount"] for m in entry["materials"]), 2)
            for m in entry["materials"]:
                key = m["material"].strip().lower()
                t = totals.setdefault(key, {"material": m["material"], "type": m["type"], "amount": 0, "days": 0})
                t["amount"] = round(t["amount"] + m["amount"], 2)
                t["days"] += 1

        events = sorted(days.values(), key=lambda e: e["date"], reverse=True)
        return jsonify({
            "status": "ok",
            "start": start, "end": end,
            "breakdown_days": len(events),
            "total_amount": round(sum(e["total_amount"] for e in events), 2),
            "totals": sorted(totals.values(), key=lambda t: -t["amount"]),
            "events": events,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/notion/gravel-summary", methods=["GET"])
@require_role("admin", "calvin")
def notion_gravel_summary():
    """
    Gravel pit sales for McCook Concrete's "Gravel Sales" tab, from the 💰 Sales
    database the gravel pit app writes to. ?start=YYYY-MM-DD&end=YYYY-MM-DD.

    - A sale's total is MATERIAL ONLY: "Total Amount" (no tax, no delivery fee),
      or tons x price per ton when the amount is blank.
    - Complete sales make up the totals and graphs. Incomplete ones are waiting
      on the customer's weight slip, so their tons/amount are estimates; they are
      counted separately.
    - Delivery: "Delivery:" in the notes = a driver was dispatched; "Customer
      Pick Up" = the customer hauled it.
    - Rows whose customer is "Test" are left out.
    """
    try:
        def _iso(name):
            v = (request.args.get(name) or "").strip()
            try:
                return datetime.date.fromisoformat(v).isoformat() if v else None
            except ValueError:
                raise ValueError(f"'{name}' must be a date like 2026-10-01")
        start, end = _iso("start"), _iso("end")

        conditions = []
        if start:
            conditions.append({"property": "Sale Date", "date": {"on_or_after": start}})
        if end:
            conditions.append({"property": "Sale Date", "date": {"on_or_before": end}})
        filter_obj = {"and": conditions} if conditions else None
        pages = notion_utils.query_data_source_all(NOTION_GRAVEL_SALES_DATASOURCE_ID,
                                                   filter_obj=filter_obj, max_pages=40)

        days, materials, customers = {}, {}, {}
        totals = {"tons": 0.0, "sales": 0.0, "loads": 0, "delivered": 0, "picked_up": 0}
        pending = {"count": 0, "est_tons": 0.0, "est_sales": 0.0}
        excluded_test = 0

        for page in pages:
            props = page.get("properties", {})
            customer = (notion_utils.prop_text(props, "Customer Name") or "").strip() or "No customer"
            if customer.lower() == "test":
                excluded_test += 1
                continue
            day = (notion_utils.prop_date(props, "Sale Date") or page.get("created_time", ""))[:10]
            if not day or (start and day < start) or (end and day > end):
                continue
            tons = notion_utils.prop_number(props, "Quantity (Tons)") or 0
            amount = notion_utils.prop_number(props, "Total Amount")
            if not amount:
                amount = tons * (notion_utils.prop_number(props, "Price Per Ton") or 0)
            material = notion_utils.prop_select(props, "Material") or "Other"

            if (notion_utils.prop_select(props, "Status") or "") == "Incomplete":
                pending["count"] += 1
                pending["est_tons"] += tons
                pending["est_sales"] += amount
                continue

            notes = notion_utils.prop_text(props, "Notes") or ""
            if "Delivery:" in notes:
                totals["delivered"] += 1
            elif "Customer Pick Up" in notes:
                totals["picked_up"] += 1

            d = days.setdefault(day, {"date": day, "tons": 0.0, "sales": 0.0, "loads": 0, "materials": {}})
            d["tons"] += tons
            d["sales"] += amount
            d["loads"] += 1
            d["materials"][material] = d["materials"].get(material, 0) + tons

            # Customer names are typed in, so "Red Willow #1" and "Red willow #1"
            # are the same customer: group case-insensitively, show the first spelling.
            for bucket, key, label in ((materials, material, material),
                                       (customers, " ".join(customer.lower().split()), customer)):
                b = bucket.setdefault(key, {"name": label, "tons": 0.0, "sales": 0.0, "loads": 0})
                b["tons"] += tons
                b["sales"] += amount
                b["loads"] += 1

            totals["tons"] += tons
            totals["sales"] += amount
            totals["loads"] += 1

        r2 = lambda x: round(x, 2)
        day_list = []
        for d in sorted(days.values(), key=lambda x: x["date"]):
            day_list.append({"date": d["date"], "tons": r2(d["tons"]), "sales": r2(d["sales"]), "loads": d["loads"],
                             "materials": {k: r2(v) for k, v in d["materials"].items()}})
        rank = lambda bucket: [dict(b, tons=r2(b["tons"]), sales=r2(b["sales"]))
                               for b in sorted(bucket.values(), key=lambda b: -b["tons"])]
        return jsonify({
            "status": "ok",
            "start": start, "end": end,
            "totals": {**totals, "tons": r2(totals["tons"]), "sales": r2(totals["sales"]),
                       "selling_days": len(day_list)},
            "days": day_list,
            "materials": rank(materials),
            "customers": rank(customers),
            "pending": {"count": pending["count"], "est_tons": r2(pending["est_tons"]),
                        "est_sales": r2(pending["est_sales"])},
            "excluded_test": excluded_test,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/notion/gravel-sales", methods=["GET"])
@require_role("admin")
def notion_gravel_sales():
    try:
        pages = notion_utils.query_data_source(NOTION_GRAVEL_SALES_DATASOURCE_ID)
        items = []
        for page in pages:
            props = page.get("properties", {})
            sale_id = notion_utils.prop_text(props, "Sale ID") or "Sale"
            customer = notion_utils.prop_text(props, "Customer Name")
            material = notion_utils.prop_select(props, "Material")
            qty = notion_utils.prop_number(props, "Quantity (Tons)")
            total = notion_utils.prop_number(props, "Total Amount")
            status = notion_utils.prop_select(props, "Status")
            invoiced = notion_utils.prop_checkbox(props, "Invoiced")
            notes = notion_utils.prop_text(props, "Notes")

            incomplete = (status == "Incomplete")
            items.append({
                "id": "notion-gravel-" + page["id"],
                "notionUrl": page.get("url"),
                "type": "general",
                "source": "notion-gravel-sales",
                "sourceLabel": "Notion · Gravel Sales",
                "job": None,
                "customer": customer,
                "title": f"{sale_id}" + (f" — {customer}" if customer else ""),
                "subtitle": f"{material or 'Material?'} · {qty or 0} tons · ${total or 0:,.0f}" + ("" if invoiced else " · not invoiced"),
                "tagClass": "warn" if incomplete else "ok",
                "tagText": status or "Unknown",
                "summary": notes or "",
                "status": "pending",
            })
        return jsonify({"count": len(items), "items": items})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Customer order lookup — searches ALL Gravel Sales history, not just
# what's currently sitting in the review queue. Available to both roles.
# ==========================================

@app.route("/api/notion/customer-search", methods=["GET"])
@require_role("admin", "calvin")
def customer_search():
    query = (request.args.get("q") or "").strip()

    try:
        filter_obj = {
            "property": "Customer Name",
            "rich_text": {"contains": query},
        } if query else None
        pages = notion_utils.query_data_source(NOTION_GRAVEL_SALES_DATASOURCE_ID, page_size=100, filter_obj=filter_obj)

        results = []
        for page in pages:
            props = page.get("properties", {})
            results.append({
                "sale_id": notion_utils.prop_text(props, "Sale ID"),
                "notion_url": page.get("url"),
                "customer": notion_utils.prop_text(props, "Customer Name"),
                "material": notion_utils.prop_select(props, "Material"),
                "quantity_tons": notion_utils.prop_number(props, "Quantity (Tons)"),
                "total_amount": notion_utils.prop_number(props, "Total Amount"),
                "sale_date": notion_utils.prop_date(props, "Sale Date"),
                "status": notion_utils.prop_select(props, "Status"),
                "invoiced": notion_utils.prop_checkbox(props, "Invoiced"),
                "delivery_address": notion_utils.prop_text(props, "Delivery Address"),
                "notes": notion_utils.prop_text(props, "Notes"),
            })

        results.sort(key=lambda r: r.get("sale_date") or "", reverse=True)
        return jsonify({"count": len(results), "results": results})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Work order lookup — search by Reporting Contact.
# NOTE: Work Orders (Service Request) lives inside a Notion multi-source
# database, which needs the newer data-sources API endpoint. This has not
# been tested against the live API from this environment — if it errors,
# check the response message and we can adjust the endpoint/version.
# ==========================================

@app.route("/api/notion/workorder-search", methods=["GET"])
@require_role("admin", "calvin")
def workorder_search():
    query = (request.args.get("q") or "").strip()

    try:
        filter_obj = {
            "property": "Reporting Contact",
            "rich_text": {"contains": query},
        } if query else None
        pages = notion_utils.query_data_source(NOTION_WORKORDERS_DATASOURCE_ID, page_size=100, filter_obj=filter_obj)

        results = []
        for page in pages:
            props = page.get("properties", {})
            results.append({
                "wo_number": notion_utils.prop_number(props, "WO #"),
                "notion_url": page.get("url"),
                "job_number": notion_utils.prop_number(props, "Job #"),
                "reporting_contact": notion_utils.prop_text(props, "Reporting Contact"),
                "work_status": notion_utils.prop_select(props, "Work Status"),
                "service_priority": notion_utils.prop_select(props, "Service Priority"),
                "date": notion_utils.prop_date(props, "Date"),
                "description_of_work": notion_utils.prop_text(props, "Description of Work"),
                "customer_complaint": notion_utils.prop_text(props, "Customer Complaint"),
            })

        results.sort(key=lambda r: r.get("date") or "", reverse=True)
        return jsonify({"count": len(results), "results": results})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Vendor lookup — search Purchase Orders by vendor, with Ordered->Delivered
# lead time computed per order and a year-over-year trend summary, so Calvin
# can see whether deliveries are taking longer than in previous years.
#
# NOTE: Vendor is a Notion "select" field, not free text, so the API can't
# do a partial-match filter on it server-side. This pulls the full PO
# history (paginated) and matches vendor names client-side instead —
# works well at this database's size, but would need a real query filter
# if the PO list grows very large.
# ==========================================

def _parse_date(d):
    if not d:
        return None
    try:
        return datetime.datetime.fromisoformat(d.replace("Z", "+00:00")).date()
    except ValueError:
        return None


@app.route("/api/notion/vendor-search", methods=["GET"])
@require_role("admin", "calvin")
def vendor_search():
    query = (request.args.get("q") or "").strip().lower()

    try:
        all_pages = notion_utils.query_data_source_all(NOTION_PURCHASE_ORDERS_DATASOURCE_ID)

        results = []
        for page in all_pages:
            props = page.get("properties", {})
            vendor = notion_utils.prop_select(props, "Vendor") or ""
            if query and query not in vendor.lower():
                continue

            ordered = notion_utils.prop_date(props, "Ordered")
            delivered = notion_utils.prop_date(props, "Delivered")

            lead_time_days = None
            ordered_d = _parse_date(ordered)
            delivered_d = _parse_date(delivered)
            if ordered_d and delivered_d:
                lead_time_days = (delivered_d - ordered_d).days

            results.append({
                "po_number": notion_utils.prop_text(props, "PO #"),
                "notion_url": page.get("url"),
                "job_number": notion_utils.prop_number(props, "Job #"),
                "vendor": vendor,
                "amount": notion_utils.prop_number(props, "Amount"),
                "status": notion_utils.prop_status(props, "Status"),
                "ordered": ordered,
                "scheduled": notion_utils.prop_date(props, "Scheduled"),
                "delivered": delivered,
                "lead_time_days": lead_time_days,
                "invoiced": notion_utils.prop_checkbox(props, "Invoiced"),
                "description": notion_utils.prop_text(props, "Description"),
            })

        results.sort(key=lambda r: r.get("ordered") or "", reverse=True)

        # Year-over-year average lead time, for spotting a "deliveries are
        # taking longer" trend. Only counts orders with both dates present.
        by_year = {}
        for r in results:
            if r["lead_time_days"] is None or not r["ordered"]:
                continue
            year = r["ordered"][:4]
            by_year.setdefault(year, []).append(r["lead_time_days"])

        yearly_trend = [
            {"year": year, "avg_lead_time_days": round(sum(days) / len(days), 1), "count": len(days)}
            for year, days in sorted(by_year.items())
        ]

        return jsonify({"count": len(results), "results": results, "yearly_trend": yearly_trend})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Fab Shop revenue — weekly file upload, parsed and stored in Dropbox.
# No live Microsoft Graph connection; you upload the workbook manually
# (admin only), the server reads it, Calvin's page just displays the result.
# ==========================================

def _find_sheet(wb, target_name):
    """Matches sheet names loosely (trailing spaces, case) since Excel tab
    names are easy to fat-finger and this shouldn't break on that."""
    target = target_name.strip().lower()
    for name in wb.sheetnames:
        if name.strip().lower() == target:
            return wb[name]
    return None


def _cell_date(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.date().isoformat() if hasattr(value, "date") else value.isoformat()
    return str(value)


def parse_fabshop_workbook(file_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    weekly = []
    weekly_sheet = _find_sheet(wb, "Weekly Labor Value")
    if weekly_sheet:
        for row in weekly_sheet.iter_rows(min_row=2, max_col=6, values_only=True):
            week_of = row[0]
            if not week_of:
                continue
            weekly.append({
                "week_of": str(week_of),
                "billable_hours": row[1] or 0,
                "non_billable_hours": row[2] or 0,
                "mgt_design_hours": row[3] or 0,
                "total_revenue": row[4] or 0,
                "total_man_hours": row[5] or 0,
            })

    daily = []
    daily_sheet = _find_sheet(wb, "Labor Value - Daily")
    if daily_sheet:
        for row in daily_sheet.iter_rows(min_row=2, max_col=6, values_only=True):
            date_val = row[0]
            if not date_val:
                continue
            daily.append({
                "date": _cell_date(date_val),
                "billable_hours": row[1] or 0,
                "non_billable_hours": row[2] or 0,
                "mgt_design_hours": row[3] or 0,
                "total_man_hours": row[4] or 0,
                "estimated_value": row[5] or 0,
            })

    return {"weekly": weekly, "daily": daily, "sheets_found": wb.sheetnames}


@app.route("/api/upload-fabshop", methods=["POST"])
@require_role("admin")
def upload_fabshop():
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "No file included in upload"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"status": "error", "message": "No file selected"}), 400

    try:
        file_bytes = file.read()
        parsed = parse_fabshop_workbook(file_bytes)

        if not parsed["weekly"] and not parsed["daily"]:
            return jsonify({
                "status": "error",
                "message": f"Couldn't find 'Weekly Labor Value' or 'Labor Value - Daily' sheets. "
                           f"Sheets found in file: {', '.join(parsed['sheets_found'])}"
            }), 400

        payload = {
            "uploaded_at": datetime.datetime.utcnow().isoformat() + "Z",
            "filename": file.filename,
            "weekly": parsed["weekly"],
            "daily": parsed["daily"],
        }

        dbx = get_dropbox_client()
        dbx.files_upload(
            json.dumps(payload, indent=2).encode("utf-8"),
            FABSHOP_REVENUE_PATH,
            mode=dropbox.files.WriteMode.overwrite,
        )

        return jsonify({
            "status": "ok",
            "weekly_rows": len(parsed["weekly"]),
            "daily_rows": len(parsed["daily"]),
            "uploaded_at": payload["uploaded_at"],
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


def _merge_fabshop_data(dbx):
    """
    Option A: the portal is now the live source of truth. This merges the
    old manually-uploaded workbook data (kept as historical backfill) with
    live daily entries submitted through the portal — live entries win for
    any date both sources have, since they're the more current record.
    Weekly rollups are computed fresh from the merged daily list (grouped
    Mon-Fri) rather than trusting the workbook's own week labels, so both
    sources roll up consistently.
    """
    legacy = _read_json_from_dropbox(dbx, FABSHOP_REVENUE_PATH, {"uploaded_at": None, "filename": None, "weekly": [], "daily": []})
    live_entries = _read_json_from_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, [])

    daily_by_date = {}
    for d in legacy.get("daily", []):
        if not d.get("date"):
            continue
        daily_by_date[d["date"]] = {
            "date": d["date"],
            "billable_hours": d.get("billable_hours", 0) or 0,
            "non_billable_hours": d.get("non_billable_hours", 0) or 0,
            "mgt_design_hours": d.get("mgt_design_hours", 0) or 0,
            "total_man_hours": d.get("total_man_hours", 0) or 0,
            "estimated_value": d.get("estimated_value", 0) or 0,
            "description": None,
            "source": "legacy_upload",
        }

    for entry in live_entries:
        totals = entry.get("totals", {})
        billable = totals.get("billable", 0) or 0
        non_billable = totals.get("non_billable", 0) or 0
        mgt_design = totals.get("mgt_design", 0) or 0
        daily_by_date[entry["date"]] = {
            "date": entry["date"],
            "billable_hours": billable,
            "non_billable_hours": non_billable,
            "mgt_design_hours": mgt_design,
            "total_man_hours": billable + non_billable + mgt_design,
            "estimated_value": entry.get("estimated_value", 0) or 0,
            "description": entry.get("description") or None,
            "source": "portal",
        }

    daily_list = sorted(daily_by_date.values(), key=lambda d: d["date"])

    weekly_by_key = {}
    for d in daily_list:
        try:
            dt = datetime.datetime.fromisoformat(d["date"]).date()
        except ValueError:
            continue
        monday = dt - datetime.timedelta(days=dt.weekday())
        friday = monday + datetime.timedelta(days=4)
        key = monday.isoformat()
        w = weekly_by_key.setdefault(key, {
            "week_of": f"{monday.strftime('%m/%d')} - {friday.strftime('%m/%d')}",
            "_sort_key": key,
            "billable_hours": 0, "non_billable_hours": 0, "mgt_design_hours": 0,
            "total_revenue": 0, "total_man_hours": 0,
        })
        w["billable_hours"] += d["billable_hours"]
        w["non_billable_hours"] += d["non_billable_hours"]
        w["mgt_design_hours"] += d["mgt_design_hours"]
        w["total_revenue"] += d["estimated_value"]
        w["total_man_hours"] += d["total_man_hours"]

    weekly_list = sorted(weekly_by_key.values(), key=lambda w: w["_sort_key"])
    for w in weekly_list:
        del w["_sort_key"]

    latest_times = [t for t in [legacy.get("uploaded_at")] + [e.get("submitted_at") for e in live_entries] if t]
    uploaded_at = max(latest_times) if latest_times else None

    return {
        "uploaded_at": uploaded_at,
        "filename": legacy.get("filename"),
        "weekly": weekly_list,
        "daily": daily_list,
    }


@app.route("/api/fabshop-revenue", methods=["GET"])
@require_role("admin", "calvin")
def fabshop_revenue():
    try:
        dbx = get_dropbox_client()
        data = _merge_fabshop_data(dbx)
        return jsonify(data)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Fab Shop portal — no login yet (by explicit request), so these routes are
# open. Revisit adding a lightweight role gate here once real usage starts.
# ==========================================

DEFAULT_FABSHOP_EMPLOYEES = ["Jon", "Jerron", "Josh", "Caleb", "Alex"]
# Old name -> the name the person goes by. Applied wherever fab shop names are shown, so
# older sheets, the stored roster and Typeform titles written as "Jason" all read "Alex".
FABSHOP_NAME_ALIASES = {"Jason": "Alex"}


def _fab_name(name):
    name = (name or "").strip()
    return FABSHOP_NAME_ALIASES.get(name, name)


def _fab_names(names):
    out = []
    for n in names or []:
        n = _fab_name(str(n))
        if n and n not in out:
            out.append(n)
    return out
# Roster names that are NOT put on a new Fab Tab Sheet automatically, only offered
# as "+ Name" buttons (e.g. Josh, who fills in now and then).
FABSHOP_OPTIONAL_PATH = os.environ.get(
    "FABSHOP_OPTIONAL_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/fabshop_employees_optional.json"
)
DEFAULT_FABSHOP_OPTIONAL = ["Josh"]
# Employment status shown on each person's row of the Fab Tab Sheet. No label for
# Jerron (shop manager). Changed from the roster manager in the portal.
FABSHOP_STATUS_PATH = os.environ.get(
    "FABSHOP_STATUS_PATH", "/Chipperfield Ag/Chipperfield/Dashboard/fabshop_employee_status.json"
)
DEFAULT_FABSHOP_STATUS = {"Jon": "Full Time", "Alex": "Part Time"}
FABSHOP_STATUS_CHOICES = ("Full Time", "Part Time")


@app.route("/api/fabshop-portal/employees", methods=["GET"])
def fabshop_portal_get_employees():
    try:
        dbx = get_dropbox_client()
        employees = _fab_names(_read_json_from_dropbox(dbx, FABSHOP_EMPLOYEES_PATH, DEFAULT_FABSHOP_EMPLOYEES))
        optional = _fab_names(_read_json_from_dropbox(dbx, FABSHOP_OPTIONAL_PATH, DEFAULT_FABSHOP_OPTIONAL))
        status = {_fab_name(k): v for k, v in _read_json_from_dropbox(dbx, FABSHOP_STATUS_PATH, DEFAULT_FABSHOP_STATUS).items()}
        return jsonify({"employees": employees, "optional": [n for n in optional if n in employees],
                        "status": {n: v for n, v in status.items() if v in FABSHOP_STATUS_CHOICES},
                        "status_choices": list(FABSHOP_STATUS_CHOICES),
                        "aliases": FABSHOP_NAME_ALIASES})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/employees", methods=["POST"])
def fabshop_portal_add_employee():
    try:
        body = request.get_json(force=True)
        name = _fab_name(body.get("name"))
        if not name:
            return jsonify({"status": "error", "message": "Name is required"}), 400

        dbx = get_dropbox_client()
        employees = _fab_names(_read_json_from_dropbox(dbx, FABSHOP_EMPLOYEES_PATH, list(DEFAULT_FABSHOP_EMPLOYEES)))
        if name not in employees:
            employees.append(name)
            _write_json_to_dropbox(dbx, FABSHOP_EMPLOYEES_PATH, employees)
        return jsonify({"status": "ok", "employees": employees})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/employees/optional", methods=["POST"])
def fabshop_portal_set_optional():
    """{"name": "Josh", "optional": true} - true: only added to a sheet when picked."""
    try:
        body = request.get_json(force=True) or {}
        name = _fab_name(body.get("name"))
        if not name:
            return jsonify({"status": "error", "message": "Name is required"}), 400
        dbx = get_dropbox_client()
        optional = [n for n in _fab_names(_read_json_from_dropbox(dbx, FABSHOP_OPTIONAL_PATH, list(DEFAULT_FABSHOP_OPTIONAL))) if n != name]
        if body.get("optional"):
            optional.append(name)
        _write_json_to_dropbox(dbx, FABSHOP_OPTIONAL_PATH, optional)
        return jsonify({"status": "ok", "optional": optional})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/employees/status", methods=["POST"])
def fabshop_portal_set_status():
    """{"name": "Alex", "status": "Part Time"} - "" removes the label."""
    try:
        body = request.get_json(force=True) or {}
        name = _fab_name(body.get("name"))
        value = (body.get("status") or "").strip()
        if not name or (value and value not in FABSHOP_STATUS_CHOICES):
            return jsonify({"status": "error", "message": "Name and Full Time / Part Time (or blank) required"}), 400
        dbx = get_dropbox_client()
        status = {_fab_name(k): v for k, v in _read_json_from_dropbox(dbx, FABSHOP_STATUS_PATH, dict(DEFAULT_FABSHOP_STATUS)).items()}
        if value:
            status[name] = value
        else:
            status.pop(name, None)
        _write_json_to_dropbox(dbx, FABSHOP_STATUS_PATH, status)
        return jsonify({"status": "ok", "employee_status": status})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/employees", methods=["DELETE"])
def fabshop_portal_remove_employee():
    try:
        body = request.get_json(force=True)
        name = _fab_name(body.get("name"))

        dbx = get_dropbox_client()
        employees = _fab_names(_read_json_from_dropbox(dbx, FABSHOP_EMPLOYEES_PATH, list(DEFAULT_FABSHOP_EMPLOYEES)))
        employees = [e for e in employees if e != name]
        _write_json_to_dropbox(dbx, FABSHOP_EMPLOYEES_PATH, employees)
        return jsonify({"status": "ok", "employees": employees})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


def _build_fabshop_entry(body, preserved_notes=None):
    """One Daily Tab entry in its stored shape. Raises ValueError if the date is missing."""
    date = body.get("date")
    if not date:
        raise ValueError("Date is required")
    employees = body.get("employees", []) or []
    totals = {"billable": 0.0, "non_billable": 0.0, "mgt_design": 0.0}
    for emp in employees:
        totals["billable"] += float(emp.get("billable") or 0)
        totals["non_billable"] += float(emp.get("non_billable") or 0)
        totals["mgt_design"] += float(emp.get("mgt_design") or 0)
    return {
        "date": date,
        "employees": employees,
        "totals": totals,
        "description": body.get("description") or "",
        "new_leads": body.get("new_leads") or "",
        "problems": body.get("problems") or "",
        "estimated_value": float(body.get("estimated_value") or 0),
        # What was written beside the dollar amount, e.g. "21774, 21764" or "covered labor (21728)"
        "value_note": (body.get("value_note") or "").strip(),
        "completed_by": body.get("completed_by") or "",
        "submitted_at": datetime.datetime.utcnow().isoformat() + "Z",
        "notes": preserved_notes or [],
        # Who was out / left early and why (used for Calvin's attendance impact)
        "absences": _clean_absences(body.get("absences")),
        "attendance_ok": [str(n) for n in (body.get("attendance_ok") or [])][:10],
        "shop_closed": bool(body.get("shop_closed")),
    }


def _clean_absences(raw):
    out = []
    for a in (raw or [])[:20]:
        if not isinstance(a, dict) or not a.get("name"):
            continue
        out.append({
            "name": str(a["name"])[:60],
            "type": "left_early" if a.get("type") == "left_early" else "out",
            "reason": str(a.get("reason") or "")[:120],
        })
    return out


@app.route("/api/fabshop-portal/daily-tab", methods=["POST"])
def fabshop_portal_submit_daily_tab():
    try:
        body = request.get_json(force=True)
        if not body.get("date"):
            return jsonify({"status": "error", "message": "Date is required"}), 400

        dbx = get_dropbox_client()
        entries = _read_json_from_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, [])

        # Preserve any existing comments/notes on this date if resubmitting
        existing = next((e for e in entries if e.get("date") == body["date"]), None)
        entry = _build_fabshop_entry(body, existing.get("notes", []) if existing else [])
        if existing:
            # An edit: keep where the sheet came from and when it was first submitted
            if existing.get("source"):
                entry["source"] = existing["source"]
            entry["first_submitted_at"] = existing.get("first_submitted_at") or existing.get("submitted_at")
            entry["edited_at"] = entry["submitted_at"]

        entries = [e for e in entries if e.get("date") != entry["date"]]  # replace same-day resubmission
        entries.append(entry)
        entries.sort(key=lambda e: e["date"])
        _write_json_to_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, entries)

        return jsonify({"status": "ok", "entry": entry})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/import-entries", methods=["POST"])
@require_role("admin")
def fabshop_portal_import_entries():
    """Admin-only bulk add of Daily Tab entries (e.g. typed up from paper
    sheets). Body: {"entries": [...same fields as a Daily Tab submit...],
    "overwrite": false}. Dates that already have an entry are skipped unless
    overwrite is true; comments on a replaced date are kept."""
    try:
        body = request.get_json(force=True) or {}
        incoming = body.get("entries")
        if not isinstance(incoming, list) or not incoming:
            return jsonify({"status": "error", "message": "Send {\"entries\": [ ... ]}"}), 400
        if len(incoming) > 400:
            return jsonify({"status": "error", "message": "At most 400 entries per import"}), 400
        overwrite = bool(body.get("overwrite"))

        dbx = get_dropbox_client()
        entries = _read_json_from_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, [])
        by_date = {e.get("date"): e for e in entries}

        added, replaced, skipped, errors = [], [], [], []
        for raw in incoming:
            try:
                d = datetime.date.fromisoformat(str(raw.get("date")))
                raw = dict(raw, date=d.isoformat())
                if raw["date"] in by_date and not overwrite:
                    skipped.append(raw["date"])
                    continue
                prev = by_date.get(raw["date"])
                entry = _build_fabshop_entry(raw, prev.get("notes", []) if prev else [])
                entry["source"] = "import"
                (replaced if prev else added).append(entry["date"])
                by_date[entry["date"]] = entry
            except Exception as err:
                errors.append(f"{raw.get('date') if isinstance(raw, dict) else raw}: {err}")

        merged = sorted(by_date.values(), key=lambda e: e["date"])
        if added or replaced:
            _write_json_to_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, merged)
        return jsonify({"status": "ok", "added": added, "replaced": replaced, "skipped": skipped, "errors": errors})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/daily-entries", methods=["GET"])
def fabshop_portal_daily_entries():
    try:
        dbx = get_dropbox_client()
        entries = _read_json_from_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, [])
        entries.sort(key=lambda e: e["date"], reverse=True)
        return jsonify({"count": len(entries), "entries": entries})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/daily-entries/<date>/notes", methods=["POST"])
def fabshop_portal_add_daily_entry_note(date):
    """Comment-style updates on a past Daily Tab entry — lets Jerron (or
    anyone) add context to a historical submission without editing it."""
    try:
        body = request.get_json(force=True)
        text = (body.get("text") or "").strip()
        if not text:
            return jsonify({"status": "error", "message": "Note text is required"}), 400

        dbx = get_dropbox_client()
        entries = _read_json_from_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, [])
        target = next((e for e in entries if e.get("date") == date), None)
        if not target:
            return jsonify({"status": "error", "message": "No entry found for that date"}), 404

        target.setdefault("notes", []).append({
            "author": body.get("author") or "Fab Shop",
            "text": text,
            "created_at": datetime.datetime.utcnow().isoformat() + "Z",
        })
        _write_json_to_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, entries)
        return jsonify({"status": "ok", "entry": target})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/leads", methods=["GET"])
def fabshop_portal_get_leads():
    try:
        dbx = get_dropbox_client()
        leads = _read_json_from_dropbox(dbx, FABSHOP_LEADS_PATH, [])
        leads.sort(key=lambda l: l.get("created_at", ""), reverse=True)
        return jsonify({"count": len(leads), "leads": leads})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/leads", methods=["POST"])
def fabshop_portal_add_lead():
    try:
        body = request.get_json(force=True)
        lead = {
            "id": "lead-" + str(int(datetime.datetime.utcnow().timestamp() * 1000)),
            "company": body.get("company") or "",
            "contact": body.get("contact") or "",
            "notes": body.get("notes") or "",
            "status": body.get("status") or "New",
            "created_at": datetime.datetime.utcnow().isoformat() + "Z",
        }
        dbx = get_dropbox_client()
        leads = _read_json_from_dropbox(dbx, FABSHOP_LEADS_PATH, [])
        leads.append(lead)
        _write_json_to_dropbox(dbx, FABSHOP_LEADS_PATH, leads)
        return jsonify({"status": "ok", "lead": lead})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/leads/<lead_id>", methods=["PATCH"])
def fabshop_portal_update_lead(lead_id):
    try:
        body = request.get_json(force=True)
        dbx = get_dropbox_client()
        leads = _read_json_from_dropbox(dbx, FABSHOP_LEADS_PATH, [])
        target = next((l for l in leads if l["id"] == lead_id), None)
        if not target:
            return jsonify({"status": "error", "message": "Lead not found"}), 404
        if "status" in body:
            target["status"] = body["status"]
        if "notes" in body:
            target["notes"] = body["notes"]
        _write_json_to_dropbox(dbx, FABSHOP_LEADS_PATH, leads)
        return jsonify({"status": "ok", "lead": target})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/leads/<lead_id>", methods=["DELETE"])
def fabshop_portal_delete_lead(lead_id):
    try:
        dbx = get_dropbox_client()
        leads = _read_json_from_dropbox(dbx, FABSHOP_LEADS_PATH, [])
        leads = [l for l in leads if l["id"] != lead_id]
        _write_json_to_dropbox(dbx, FABSHOP_LEADS_PATH, leads)
        return jsonify({"status": "ok"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-portal/evaluation-data", methods=["GET"])
def fabshop_portal_evaluation_data():
    """Same merged data as /api/fabshop-revenue, but open (no login) since
    the portal itself has no auth yet — used for the Evaluation graph."""
    try:
        dbx = get_dropbox_client()
        data = _merge_fabshop_data(dbx)
        return jsonify(data)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


FABSHOP_TRACKING_CHANGE_WEEK = "2025-11-03"  # before this, non-billable hours were barely recorded


@app.route("/api/fabshop-portal/trends", methods=["GET"])
def fabshop_portal_trends():
    """Hours trends for the Evaluation tab: one row per Monday-week (and per
    month) from the first day on file to the last. A week/month with no sheet
    is `has_data: false` so charts break instead of dropping to zero. Open
    like evaluation-data, since the portal has no login yet."""
    try:
        dbx = get_dropbox_client()
        daily = [d for d in _merge_fabshop_data(dbx)["daily"] if d.get("date")]
        if not daily:
            return jsonify({"has_data": False, "weeks": [], "months": []})
        live = {e["date"]: e for e in _read_json_from_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, []) if e.get("date")}
        outside_re = [pat for pat, label, kind in FABSHOP_ELSEWHERE_PATTERNS]

        def blank(key, label):
            return {"key": key, "label": label, "has_data": False, "days": 0, "billable": 0.0, "non_billable": 0.0,
                    "mgt_design": 0.0, "total": 0.0, "person_days": 0, "outside_dates": [], "value": 0.0}

        def add(row, d):
            b, nb, m = (float(d.get(k) or 0) for k in ("billable_hours", "non_billable_hours", "mgt_design_hours"))
            if b + nb + m <= 0:
                return  # a day with no hours (shop closed, holiday) is not a data point
            row["has_data"] = True
            row["days"] += 1
            row["billable"] += b; row["non_billable"] += nb; row["mgt_design"] += m
            row["total"] += b + nb + m
            row["value"] += float(d.get("estimated_value") or 0)
            e = live.get(d["date"])
            if e:
                row["person_days"] += sum(1 for h in _entry_hours_by_person(e).values() if h > 0)
                text = f"{e.get('description') or ''} {e.get('value_note') or ''}"
                if any(p.search(text) for p in outside_re):
                    row["outside_dates"].append(d["date"])

        first = datetime.date.fromisoformat(_monday_of(daily[0]["date"]))
        last = datetime.date.fromisoformat(_monday_of(daily[-1]["date"]))
        weeks, w = {}, first
        while w <= last:
            weeks[w.isoformat()] = blank(w.isoformat(), f"{w.month}/{w.day}")
            w += datetime.timedelta(weeks=1)
        months = {}
        for d in daily:
            add(weeks[_monday_of(d["date"])], d)
            mk = d["date"][:7]
            months.setdefault(mk, blank(mk, datetime.date.fromisoformat(d["date"]).strftime("%b %Y")))
            add(months[mk], d)

        def finish(r):
            r = dict(r)
            for k in ("billable", "non_billable", "mgt_design", "total", "value"):
                r[k] = round(r[k], 2)
            if r["has_data"]:
                r["billable_pct"] = round(r["billable"] / r["total"] * 100, 1)
                r["billable_per_person_day"] = round(r["billable"] / r["person_days"], 2) if r["person_days"] else None
            else:
                r["billable_pct"] = r["billable_per_person_day"] = None
            return r

        wk = [finish(r) for r in weeks.values()]
        mo = [finish(months[k]) for k in sorted(months)]
        return jsonify({
            "has_data": True, "weeks": wk, "months": mo,
            "target_pct": 80,
            "tracking_change_week": FABSHOP_TRACKING_CHANGE_WEEK,
            "first_date": daily[0]["date"], "last_date": daily[-1]["date"],
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ---------------------------------------------------------------------------
# Calvin's Fab Shop screen: Summary (line graph) + Weekly Breakdown (day cards)
# ---------------------------------------------------------------------------

# Places named in a Daily Tab's notes that mean hours went somewhere else.
FABSHOP_ELSEWHERE_PATTERNS = [
    (re.compile(r"\bmci\b|mc\s?cook concrete|cement plant|concrete plant", re.IGNORECASE), "McCook Concrete (MCI)", "mci"),
    (re.compile(r"concrete truck|drove concrete|drive concrete|driving concrete", re.IGNORECASE), "McCook Concrete (MCI)", "mci"),
    (re.compile(r"gravel pit", re.IGNORECASE), "Gravel pit", "other"),
]


# ---- Holidays, attendance and impact ---------------------------------------
FABSHOP_COVERAGE_START = "2025-08-18"   # per-person hours begin here
FABSHOP_JOSH_LAST_DAY = "2026-02-27"
FABSHOP_PART_TIMERS = {"Caleb": "Caleb W. (part-time)"}


def _display_name(name):
    name = _fab_name(name)
    return FABSHOP_PART_TIMERS.get(name, name)


def _holiday_name(d):
    """Shop holiday for a date, or ''. Observed rules for the US holidays the
    shop takes, plus Dec 24/26/31 which it also closes for."""
    if isinstance(d, str):
        d = datetime.date.fromisoformat(d)
    y, m, day = d.year, d.month, d.day
    if m == 1 and day == 1:
        return "New Year's Day"
    if m == 5 and d.weekday() == 0 and day + 7 > 31:
        return "Memorial Day"
    jul4 = datetime.date(y, 7, 4)
    obs = jul4 - datetime.timedelta(days=1) if jul4.weekday() == 5 else (
        jul4 + datetime.timedelta(days=1) if jul4.weekday() == 6 else jul4)
    if d == obs:
        return "Independence Day (observed)" if obs != jul4 else "Independence Day"
    if m == 9 and d.weekday() == 0 and day <= 7:
        return "Labor Day"
    if m == 11 and d.weekday() == 3 and 22 <= day <= 28:
        return "Thanksgiving"
    if m == 11 and d.weekday() == 4 and 23 <= day <= 29:
        return "Day after Thanksgiving"
    if m == 12 and day in (24, 25, 26):
        return {24: "Christmas Eve", 25: "Christmas Day", 26: "Day after Christmas"}[day]
    if m == 12 and day == 31:
        return "New Year's Eve"
    return ""


def _core_roster(date_str):
    """Full-time crew expected on a date (part-timers are never 'absent')."""
    if date_str < FABSHOP_COVERAGE_START:
        return []
    crew = ["Jon", "Jerron"]
    if "2025-09-22" <= date_str <= FABSHOP_JOSH_LAST_DAY:
        crew.append("Josh")
    return crew


def _entry_hours_by_person(entry):
    return {e.get("name"): sum(float(e.get(k) or 0) for k in ("billable", "non_billable", "mgt_design"))
            for e in entry.get("employees", []) or [] if e.get("name")}


def _entry_absences(entry):
    """Explicit absences plus implied ones (core crew with no hours and no
    explanation). Returns (list, shop_closed)."""
    date = entry["date"]
    hours = _entry_hours_by_person(entry)
    total = sum(hours.values())
    explicit = [dict(a) for a in (entry.get("absences") or [])]
    if entry.get("shop_closed") or any(a["name"] == "Entire fab shop" for a in explicit):
        return [a for a in explicit if a["name"] == "Entire fab shop"] or \
            [{"name": "Entire fab shop", "type": "out", "reason": "everyone gone"}], True
    named = {a["name"] for a in explicit}
    ok = set(entry.get("attendance_ok") or [])
    if total <= 0 and _holiday_name(date):
        return [], False
    if total <= 0 and date >= FABSHOP_COVERAGE_START:
        return [{"name": "Entire fab shop", "type": "out", "reason": "no hours logged"}], True
    for n in _core_roster(date):
        if n not in named and n not in ok and hours.get(n, 0) <= 0:
            explicit.append({"name": n, "type": "out", "reason": "not listed"})
    return explicit, False


def _reason_category(reason):
    r = (reason or "").lower()
    if "holiday" in r:
        return "Holiday time off"
    if "sick" in r or "not feeling" in r or "ill" in r.split():
        return "Sick"
    if "pto" in r:
        return "PTO"
    if "doctor" in r or "appointment" in r:
        return "Appointment"
    if "family" in r or "wife" in r or "son" in r.split():
        return "Family"
    if "half day" in r or "short" in r:
        return "Short day"
    if r in ("", "not listed", "no hours logged"):
        return "No reason written"
    return "Other (gone / off)"


def _weekdays_between(a, b):
    d = datetime.date.fromisoformat(a)
    end = datetime.date.fromisoformat(b)
    while d <= end:
        if d.weekday() < 5:
            yield d.isoformat()
        d += datetime.timedelta(days=1)


def _hm(hours):
    """Decimal hours -> 'H:MM' (rounded to the minute)."""
    minutes = int(round(float(hours or 0) * 60))
    return f"{minutes // 60}:{minutes % 60:02d}"


def _pct_split(billable, non_billable, mgt):
    total = billable + non_billable + mgt
    if total <= 0:
        return {"billable": 0.0, "non_billable": 0.0, "mgt_design": 0.0}
    return {
        "billable": round(billable / total * 100, 1),
        "non_billable": round(non_billable / total * 100, 1),
        "mgt_design": round(mgt / total * 100, 1),
    }


def _monday_of(date_str):
    d = datetime.date.fromisoformat(date_str)
    return (d - datetime.timedelta(days=d.weekday())).isoformat()


def _fabshop_day_card(day, live_entry):
    """One day for the Weekly Breakdown. `day` is a row from
    _merge_fabshop_data()['daily']; `live_entry` is the portal/import entry
    for that date if there is one (it carries who did what, notes, leads)."""
    b = float(day.get("billable_hours") or 0)
    nb = float(day.get("non_billable_hours") or 0)
    m = float(day.get("mgt_design_hours") or 0)
    value = float(day.get("estimated_value") or 0)
    card = {
        "date": day["date"],
        # portal = online form, import = typed in from a paper sheet, excel = old workbook upload
        "source": (live_entry.get("source") or "portal") if live_entry else "excel",
        "total_hm": _hm(b + nb + m),
        "billable_hm": _hm(b), "non_billable_hm": _hm(nb), "mgt_design_hm": _hm(m),
        "billable": b, "non_billable": nb, "mgt_design": m,
        "split": _pct_split(b, nb, m),
        "value": value,
        "value_pending": value <= 0,
        "value_note": "",
        "wo_tags": [],
        "elsewhere": [],
        "people": [],
        "description": "",
        "new_leads": "",
        "problems": "",
        "completed_by": "",
        "notes": [],
        "holiday": _holiday_name(day["date"]),
        "absences": [],
        "shop_closed": False,
        "no_notes": True,
    }
    if live_entry:
        note = (live_entry.get("value_note") or "").strip()
        text = f"{live_entry.get('description') or ''} {note}"
        people = []
        for emp in live_entry.get("employees", []) or []:
            tot = sum(float(emp.get(k) or 0) for k in ("billable", "non_billable", "mgt_design"))
            if emp.get("name"):
                people.append({"name": _display_name(emp["name"]), "total_hm": _hm(tot)})
        abs_list, closed = _entry_absences(dict(live_entry, date=day["date"]))
        card.update({
            "value_note": note,
            "wo_tags": live_entry.get("wo_tags") or djl_titles.id_tags(note),
            "elsewhere": list({label: {"label": label, "kind": kind}
                               for pat, label, kind in FABSHOP_ELSEWHERE_PATTERNS if pat.search(text)}.values()),
            "people": people,
            "description": live_entry.get("description") or "",
            "new_leads": live_entry.get("new_leads") or "",
            "problems": live_entry.get("problems") or "",
            "completed_by": live_entry.get("completed_by") or "",
            "edited_at": live_entry.get("edited_at") or "",
            "notes": live_entry.get("notes") or [],
            "absences": [dict(a, name=_display_name(a["name"])) for a in abs_list],
            "shop_closed": closed,
            "no_notes": not (live_entry.get("description") or "").strip(),
        })
    return card


@app.route("/api/fabshop-weekly", methods=["GET"])
@require_role("admin", "calvin")
def fabshop_weekly():
    """One week of day cards (newest first) plus the week's totals.
    ?week=YYYY-MM-DD (any date; snapped to its Monday). Default: latest week with data."""
    try:
        dbx = get_dropbox_client()
        merged = _merge_fabshop_data(dbx)
        daily = [d for d in merged["daily"] if d.get("date")]
        if not daily:
            return jsonify({"has_data": False})

        live_by_date = {e.get("date"): e for e in _read_json_from_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, [])}
        mondays = sorted({_monday_of(d["date"]) for d in daily})

        week = mondays[-1]
        requested = request.args.get("week")
        if requested:
            try:
                snapped = _monday_of(requested)
                if snapped in mondays:
                    week = snapped
            except ValueError:
                pass

        i = mondays.index(week)
        days = [d for d in daily if _monday_of(d["date"]) == week]
        cards = sorted((_fabshop_day_card(d, live_by_date.get(d["date"])) for d in days),
                       key=lambda c: c["date"], reverse=True)

        b = sum(c["billable"] for c in cards)
        nb = sum(c["non_billable"] for c in cards)
        m = sum(c["mgt_design"] for c in cards)
        start = datetime.date.fromisoformat(week)
        end = start + datetime.timedelta(days=4)
        if any(datetime.date.fromisoformat(c["date"]).weekday() >= 5 for c in cards):
            end = start + datetime.timedelta(days=6)

        present = {c["date"] for c in cards}
        span_lo, span_hi = daily[0]["date"], daily[-1]["date"]
        off_days = []
        for wd in _weekdays_between(week, (start + datetime.timedelta(days=4)).isoformat()):
            if wd not in present and span_lo <= wd <= span_hi:
                hol = _holiday_name(wd)
                off_days.append({"date": wd, "holiday": hol,
                                 "label": f"Holiday: {hol}" if hol else "No tab on file"})
        off_days.sort(key=lambda o: o["date"], reverse=True)

        return jsonify({
            "has_data": True,
            "week": week,
            "off_days": off_days,
            "week_label": f"{start.strftime('%b')} {start.day} \u2013 {end.strftime('%b')} {end.day}, {end.year}",
            "prev_week": mondays[i - 1] if i > 0 else None,
            "next_week": mondays[i + 1] if i < len(mondays) - 1 else None,
            "summary": {
                "total_hm": _hm(b + nb + m),
                "billable_hm": _hm(b), "non_billable_hm": _hm(nb), "mgt_design_hm": _hm(m),
                "split": _pct_split(b, nb, m),
                "value": sum(c["value"] for c in cards),
                "pending_days": sum(1 for c in cards if c["value_pending"]),
            },
            "days": cards,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-summary", methods=["GET"])
@require_role("admin", "calvin")
def fabshop_summary():
    """Line-graph data. ?range=week (the latest week with data, one point per
    day) or ?range=6m (26 weeks ending at the latest week with data, one
    point per week). A day/week with no data is null (a break in the line),
    never zero."""
    try:
        rng = request.args.get("range", "week")
        if rng not in ("week", "6m"):
            return jsonify({"status": "error", "message": "range must be week or 6m"}), 400

        dbx = get_dropbox_client()
        daily = [d for d in _merge_fabshop_data(dbx)["daily"] if d.get("date")]
        if not daily:
            return jsonify({"has_data": False, "range": rng, "points": []})

        by_date = {d["date"]: d for d in daily}
        latest_week = datetime.date.fromisoformat(_monday_of(daily[-1]["date"]))
        first_date = daily[0]["date"]
        last_date = daily[-1]["date"]

        def point(label, key, rows):
            if not rows:
                return {"label": label, "key": key, "pending": False, "value": None, "billable": None, "total_hours": None}
            val = round(sum(float(r.get("estimated_value") or 0) for r in rows), 2)
            return {
                "label": label, "key": key,
                # Hours on file but no dollar value yet: not $0, so leave a break in the line
                "pending": val <= 0,
                "value": val if val > 0 else None,
                "billable": round(sum(float(r.get("billable_hours") or 0) for r in rows), 2),
                "total_hours": round(sum(float(r.get("total_man_hours") or 0) for r in rows), 2),
            }

        points, gaps = [], []
        if rng == "week":
            for n in range(5):
                d = latest_week + datetime.timedelta(days=n)
                row = by_date.get(d.isoformat())
                points.append(point(d.strftime("%a ") + f"{d.month}/{d.day}", d.isoformat(), [row] if row else []))
        else:
            weeks = [latest_week - datetime.timedelta(weeks=n) for n in range(25, -1, -1)]
            by_week = {}
            for d in daily:
                by_week.setdefault(_monday_of(d["date"]), []).append(d)
            for w in weeks:
                points.append(point(f"{w.month}/{w.day}", w.isoformat(), by_week.get(w.isoformat(), [])))
            run_start = None
            for idx, pt in enumerate(points + [{"value": 1, "key": None}]):
                if pt["value"] is None:
                    run_start = idx if run_start is None else run_start
                elif run_start is not None:
                    if idx - run_start >= 4:  # a month or more with nothing on file
                        gaps.append({"from": points[run_start]["key"], "to": points[idx - 1]["key"], "weeks": idx - run_start})
                    run_start = None

        have = [p for p in points if p["value"] is not None]
        return jsonify({
            "has_data": True, "range": rng, "points": points, "gaps": gaps,
            "first_data_date": first_date, "last_data_date": last_date,
            "total_value": round(sum(p["value"] for p in have), 2),
            "periods_with_data": len(have),
            "average_value": round(sum(p["value"] for p in have) / len(have), 2) if have else None,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-impact", methods=["GET"])
@require_role("admin", "calvin")
def fabshop_impact():
    """Holiday + absenteeism impact for Calvin. Hours/value lost are
    ESTIMATES: each absent person's usual full-day hours, priced at the
    shop's average value per billable hour."""
    try:
        dbx = get_dropbox_client()
        merged = _merge_fabshop_data(dbx)
        daily = [d for d in merged["daily"] if d.get("date")]
        live = {e["date"]: e for e in _read_json_from_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, []) if e.get("date")}
        if not live:
            return jsonify({"has_data": False})

        # Usual full-day hours per person (days of 6h+ only)
        full = {}
        for e in live.values():
            if e["date"] < FABSHOP_COVERAGE_START:
                continue
            for n, h in _entry_hours_by_person(e).items():
                if h >= 6:
                    full.setdefault(n, []).append(h)
        usual = {n: sum(v) / len(v) for n, v in full.items() if v}

        # $ per hour of output on days that have a value
        val_days = [d for d in daily if (d.get("estimated_value") or 0) > 0 and d["date"] >= FABSHOP_COVERAGE_START]
        tot_hours = sum(d["total_man_hours"] for d in val_days)
        tot_value = sum(d["estimated_value"] for d in val_days)
        value_per_hour = (tot_value / tot_hours) if tot_hours > 0 else 0.0

        events, by_person, by_cat = [], {}, {}
        totals = {"out": 0, "left_early": 0}
        closed_days, hours_lost_total = [], 0.0
        for date in sorted(live):
            if date < FABSHOP_COVERAGE_START:
                continue
            e = live[date]
            hrs = _entry_hours_by_person(e)
            abs_list, closed = _entry_absences(e)
            if closed:
                lost = sum(usual.get(n, 0) for n in _core_roster(date))
                note = (e.get("description") or "").strip()
                closed_days.append({"date": date, "note": note, "est_hours": round(lost, 1)})
                hours_lost_total += lost
                continue
            for a in abs_list:
                n = a["name"]
                if a["type"] == "left_early":
                    lost = max(0.0, usual.get(n, 0) - hrs.get(n, 0))
                    totals["left_early"] += 1
                else:
                    lost = usual.get(n, 0)
                    totals["out"] += 1
                hours_lost_total += lost
                cat = _reason_category(a["reason"])
                c = by_cat.setdefault(cat, {"reason": cat, "days": 0, "est_hours": 0.0})
                c["days"] += 1; c["est_hours"] += lost
                pp = by_person.setdefault(n, {"name": _display_name(n), "out": 0, "left_early": 0, "est_hours": 0.0})
                pp[a["type"]] += 1; pp["est_hours"] += lost
                events.append({"date": date, "name": _display_name(n), "type": a["type"],
                               "reason": a["reason"], "category": cat, "est_hours": round(lost, 1)})

        # Holidays: weekdays the shop took off with no hours on file
        last = daily[-1]["date"]
        hol_days, hol_hours = [], 0.0
        for wd in _weekdays_between(FABSHOP_COVERAGE_START, last):
            name = _holiday_name(wd)
            if not name:
                continue
            e = live.get(wd)
            worked = e is not None and sum(_entry_hours_by_person(e).values()) > 0
            lost = 0.0 if worked else sum(usual.get(n, 0) for n in _core_roster(wd))
            hol_days.append({"date": wd, "name": name, "worked": worked, "est_hours": round(lost, 1)})
            hol_hours += lost

        have = {d["date"] for d in daily}
        no_tab = [wd for wd in _weekdays_between(FABSHOP_COVERAGE_START, last)
                  if wd not in have and not _holiday_name(wd)]

        def money(h):
            return round(h * value_per_hour)

        return jsonify({
            "has_data": True,
            "coverage_from": FABSHOP_COVERAGE_START, "through": last,
            "estimate_note": "Estimates: each person's usual full-day hours, priced at the shop's average value per hour on days with a value.",
            "value_per_hour": round(value_per_hour, 2),
            "absences": {
                "person_days_out": totals["out"], "person_days_left_early": totals["left_early"],
                "est_hours": round(sum(c["est_hours"] for c in by_cat.values()), 1),
                "est_value": money(sum(c["est_hours"] for c in by_cat.values())),
                "by_reason": sorted(({"reason": c["reason"], "days": c["days"], "est_hours": round(c["est_hours"], 1)}
                                     for c in by_cat.values()), key=lambda c: -c["days"]),
                "by_person": sorted(({"name": p["name"], "out": p["out"], "left_early": p["left_early"],
                                      "est_hours": round(p["est_hours"], 1)} for p in by_person.values()),
                                    key=lambda p: -p["out"]),
            },
            "shop_closed": {"count": len(closed_days), "days": closed_days,
                            "est_hours": round(sum(c["est_hours"] for c in closed_days), 1),
                            "est_value": money(sum(c["est_hours"] for c in closed_days))},
            "holidays": {"count": sum(1 for h in hol_days if not h["worked"]), "days": hol_days,
                         "est_hours": round(hol_hours, 1), "est_value": money(hol_hours)},
            "no_tab_days": {"count": len(no_tab), "note": "Weekdays with no sheet on file. Not counted as absences."},
            "total_est_hours": round(hours_lost_total + hol_hours, 1),
            "total_est_value": money(hours_lost_total + hol_hours),
            "recent_events": sorted(events, key=lambda x: x["date"], reverse=True)[:25],
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# The fab shop's dedicated welders. They also fill in Daily Job Logs for jobs, and
# those logs show on the Fab Shop pages too (matched on the first name in the title).
# "Jason" stays in the list: Alex's older Daily Job Logs are titled with that name.
FABSHOP_WELDERS = {n.strip().lower() for n in os.environ.get("FABSHOP_WELDERS", "Jon,Alex,Jason").split(",") if n.strip()}


def _is_welder(name):
    first = (name or "").strip().split(" ")[0].lower()
    return first in FABSHOP_WELDERS


def _is_fabshop_log(log):
    """Logs that belong on the Fab Shop pages: the shop's daily team log, Jerron's
    Daily Management Log, and any Daily Job Log by one of the shop's welders."""
    kind = log.get("job_or_wo") or ""
    name = log.get("name") or ""
    return kind == "Fab Shop Daily Log" or (
        kind == "Daily Management Log" and "jerron" in name.lower()) or _is_welder(name)


def _date_arg(name):
    raw = (request.args.get(name) or "").strip()
    if not raw:
        return None
    return datetime.date.fromisoformat(raw)


@app.route("/api/fabshop-day-log", methods=["GET"])
@require_role("admin", "calvin")
def fabshop_day_log():
    """
    Everything the Fab Shop logged, grouped by day (newest first):
      tab  - the Daily Tab from the online form (hours, value, notes), or older
             spreadsheet history, as one Weekly Breakdown day card
      logs - Typeform logs in Dropbox (Fab Shop Daily Log, Jerron's management log)
    ?tabs_only=1 skips the Dropbox logs (the portal calendar shows tab sheets only).
    ?start=YYYY-MM-DD&end=YYYY-MM-DD (default: the last 7 days). Any length works
    (Jerron's History has an "All time" view): only Fab Shop logs have their text read.
    """
    try:
        try:
            start, end = _date_arg("start"), _date_arg("end")
        except ValueError:
            return jsonify({"status": "error", "message": "Dates must look like 2026-10-29"}), 400
        today = _today_central()
        end = end or today
        start = start or (end - datetime.timedelta(days=6))
        if start > end:
            start, end = end, start
        start = max(start, datetime.date(2020, 1, 1))
        lo, hi = start.isoformat(), end.isoformat()
        span = (end - start).days + 1

        dbx = get_dropbox_client()
        live = {e.get("date"): e for e in _read_json_from_dropbox(dbx, FABSHOP_DAILY_ENTRIES_PATH, []) if e.get("date")}
        days = {}
        for d in _merge_fabshop_data(dbx)["daily"]:
            if d.get("date") and lo <= d["date"] <= hi:
                days[d["date"]] = {"date": d["date"], "tab": _fabshop_day_card(d, live.get(d["date"])), "logs": []}

        logs_error = None
        try:
            if request.args.get("tabs_only") == "1":
                raise LookupError("skipped")
            logs = [l for l in _apply_log_corrections(
                        dbx, get_recent_daily_logs(start=start, end=end, keep=_is_fabshop_log, max_days=span))
                    if _is_fabshop_log(l)]
        except LookupError:
            logs = []
        except Exception as err:          # the Daily Tabs still show if Dropbox logs can't be read
            logs, logs_error = [], str(err)
        for l in logs:
            first, _, rest = (l.get("name") or "").partition(" ")
            if first in FABSHOP_NAME_ALIASES:
                l["name"] = (FABSHOP_NAME_ALIASES[first] + (" " + rest if rest else "")).strip()
            if lo <= (l.get("date") or "") <= hi:
                days.setdefault(l["date"], {"date": l["date"], "tab": None, "logs": []})["logs"].append(l)

        out = sorted(days.values(), key=lambda x: x["date"], reverse=True)
        for day in out:
            day["holiday"] = _holiday_name(day["date"])
        return jsonify({
            "status": "ok", "start": lo, "end": hi, "days": out,
            "tab_count": sum(1 for d in out if d["tab"]),
            "log_count": sum(len(d["logs"]) for d in out),
            "logs_error": logs_error,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/fabshop-pto", methods=["GET"])
@require_role("admin", "calvin")
def fabshop_pto():
    """PTO / time off / sick time for the fab shop crew, from the same Notion
    request forms as the Chipperfield calendar. Matched by first name against the
    fab shop employee list. Upcoming (next 90 days) and recent (last 30 days)."""
    try:
        dbx = get_dropbox_client()
        roster = _read_json_from_dropbox(dbx, FABSHOP_EMPLOYEES_PATH, list(DEFAULT_FABSHOP_EMPLOYEES)) or []
        roster = _fab_names(roster)
        first_names = {str(n).strip().split()[0].lower() for n in roster if str(n).strip()}
        first_names |= {old.lower() for old, new in FABSHOP_NAME_ALIASES.items() if new.lower() in first_names}
        today = _today_central()
        lo = (today - datetime.timedelta(days=30)).isoformat()
        hi = (today + datetime.timedelta(days=90)).isoformat()
        items = []
        for r in _pto_requests():
            first = (r["name"].split() or [""])[0].lower()
            if first not in first_names or r["end"] < lo or r["start"] > hi:
                continue
            items.append({**r, "label": _pto_label(r["types"]),
                          "when": "now" if r["start"] <= today.isoformat() <= r["end"]
                          else ("upcoming" if r["start"] > today.isoformat() else "past")})
        items.sort(key=lambda r: r["start"])
        return jsonify({"status": "ok", "today": today.isoformat(), "roster": roster, "items": items})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ---------------------------------------------------------------------------
# Jerron's portal: the same Fab Shop views as Calvin's page. The portal has no
# login (leadership's call), so these are the Calvin endpoints without the role
# check, and file links are limited to Fab Shop logs.
# ---------------------------------------------------------------------------
for _rule, _view in (("summary", fabshop_summary), ("weekly", fabshop_weekly), ("impact", fabshop_impact),
                     ("day-log", fabshop_day_log), ("pto", fabshop_pto)):
    app.add_url_rule(f"/api/fabshop-portal/{_rule}", f"fabshop_portal_{_rule.replace('-', '_')}",
                     _view.__wrapped__, methods=["GET"])


def _is_fabshop_log_path(path):
    """True for a Fab Shop Daily Log file, or Jerron's Daily Management Log, in the job log folders."""
    if not isinstance(path, str) or ".." in path or not path.startswith(DROPBOX_BASE_FOLDER + "/"):
        return False
    parts = [p.strip() for p in path.rsplit("/", 1)[-1].rsplit(".", 1)[0].split("|")]
    if "Fab Shop Daily Log" in parts:
        return True
    if len(parts) > 1 and _is_welder(parts[1]):
        return True
    return "Daily Management Log" in parts and len(parts) > 1 and "jerron" in parts[1].lower()


@app.route("/api/fabshop-portal/log-links", methods=["POST"])
def fabshop_portal_log_links():
    """Same as /api/log-links, for Jerron's portal: Fab Shop logs only."""
    try:
        body = request.get_json(force=True) or {}
        paths = body.get("paths") or []
        mode = body.get("mode", "view")
        if mode not in ("view", "download", "temp"):
            return jsonify({"status": "error", "message": "Bad mode"}), 400
        if not paths or len(paths) > 30:
            return jsonify({"status": "error", "message": "Between 1 and 30 paths required"}), 400
        if not all(_is_fabshop_log_path(p) for p in paths):
            return jsonify({"status": "error", "message": "Only Fab Shop log files can be opened here"}), 403
        return jsonify({"status": "ok", "links": get_file_links(get_dropbox_client(), paths, mode)})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/shared/<path:name>")
def shared_asset(name):
    """Styles and scripts shared by Calvin's Fab Shop page and Jerron's portal."""
    return send_from_directory(os.path.join(_REPO_ROOT, "shared"), name)


@app.route("/api/fabshop-portal/jerron-logs", methods=["GET"])
def fabshop_portal_jerron_logs():
    try:
        days = int(request.args.get("days", 30))
        logs = get_recent_daily_logs(days=days)
        jerron_logs = [
            l for l in logs
            if "jerron" in (l.get("name") or "").lower()
            and l.get("job_or_wo") in ("Fab Shop Daily Log", "Daily Management Log")
        ]
        return jsonify({"count": len(jerron_logs), "logs": jerron_logs})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Current Jobs — the same data reviewed in Monday morning meetings.
# NOTE: like Work Orders, this is an inline database and uses the newer
# data-sources endpoint. Please test and report back if it errors.
# ==========================================

@app.route("/api/notion/current-jobs", methods=["GET"])
@require_role("admin", "calvin")
def notion_current_jobs():
    try:
        pages = notion_utils.query_data_source(NOTION_CURRENT_JOBS_DATASOURCE_ID, page_size=100)
        items = []
        for page in pages:
            props = page.get("properties", {})
            items.append({
                "job_number": notion_utils.prop_text(props, "Job #"),
                "notion_url": page.get("url"),
                "name": notion_utils.prop_text(props, "Name"),
                "job_status": notion_utils.prop_select(props, "Job Status"),
                "stage": notion_utils.prop_select(props, "Stage"),
                "priority": notion_utils.prop_select(props, "Priority"),
                "project_manager": notion_utils.prop_select(props, "Project Manager"),
                "hours_budget": notion_utils.prop_number(props, "Hours budget"),
                "hours_used": notion_utils.prop_number(props, "Hours Used"),
                "total_budget": notion_utils.prop_number(props, "Total Budget"),
                "contract": notion_utils.prop_number(props, "Contract"),
                "next_action": notion_utils.prop_text(props, "Next Action"),
                "pm_next_action": notion_utils.prop_text(props, "PM Next Action"),
                "status_note": notion_utils.prop_text(props, "Status Note"),
                "calvin_review": notion_utils.prop_checkbox(props, "Calvin Review"),
            })

        # Active-feeling jobs first: anything not Won/Lost/Cancelled/Complete-ish
        closed_statuses = {"Won", "Lost", "Cancelled", "Complete", "Expired"}
        items.sort(key=lambda i: (i["job_status"] in closed_statuses if i["job_status"] else False, i["job_number"] or ""))

        return jsonify({"count": len(items), "items": items})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Job budget links — reads the "Current Job Analyses" Notion PAGE (not a
# database), which is a running list of headings (week-of dates) each
# followed by links to per-job SharePoint budget workbooks. Filtered to
# the current calendar month only, since jobs repeat across many weeks.
#
# NOTE: this parses raw Notion page blocks rather than a structured
# database, which is inherently less predictable than the other routes.
# Please test and report back what it returns.
# ==========================================

_DATE_IN_HEADING = re.compile(r'(\d{1,2})/(\d{1,2})/(\d{2,4})')
_JOB_NUMBER_IN_TEXT = re.compile(r'J\d{3,5}', re.IGNORECASE)


def _parse_job_budget_links():
    blocks = notion_utils.get_page_blocks(NOTION_JOB_ANALYSES_PAGE_ID)
    now = datetime.datetime.utcnow().date()
    current_heading_date = None
    by_job = {}

    for block in blocks:
        btype = block.get("type")

        if btype in ("heading_1", "heading_2", "heading_3"):
            text, _ = notion_utils.block_plain_text(block)
            m = _DATE_IN_HEADING.search(text)
            if m:
                mm, dd, yy = m.groups()
                yy = int(yy)
                if yy < 100:
                    yy += 2000
                try:
                    current_heading_date = datetime.date(yy, int(mm), int(dd))
                except ValueError:
                    current_heading_date = None
            else:
                current_heading_date = None  # a non-date heading, e.g. "Purchase Order Review"
            continue

        if not current_heading_date:
            continue
        if current_heading_date.year != now.year or current_heading_date.month != now.month:
            continue

        text, url = notion_utils.block_plain_text(block)
        if not url or not text:
            continue
        job_match = _JOB_NUMBER_IN_TEXT.search(text)
        if not job_match:
            continue

        job_number = job_match.group(0).upper()
        by_job.setdefault(job_number, []).append({
            "date": current_heading_date.isoformat(),
            "filename": text,
            "url": url,
        })

    return by_job


@app.route("/api/notion/job-budgets", methods=["GET"])
@require_role("admin", "calvin")
def job_budgets():
    try:
        data = _parse_job_budget_links()
        return jsonify({"count": len(data), "budgets_by_job": data})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


def _parse_latest_job_analyses():
    """
    Current Jobs, redefined: not the full Current Jobs database (which
    includes leads, prospects, and everything else), but specifically
    whichever job numbers appear under the MOST RECENT dated heading on
    the Current Job Analyses page — that's what's actually being actively
    tracked week to week, updated every Friday/Monday.
    """
    blocks = notion_utils.get_page_blocks(NOTION_JOB_ANALYSES_PAGE_ID)
    as_of_date = None
    as_of_label = None
    jobs = []
    found_latest_heading = False

    for block in blocks:
        btype = block.get("type")

        if btype in ("heading_1", "heading_2", "heading_3"):
            text, _ = notion_utils.block_plain_text(block)
            m = _DATE_IN_HEADING.search(text)
            if m:
                if found_latest_heading:
                    break  # hit the *next* dated heading — stop, we only want the first (latest) one
                mm, dd, yy = m.groups()
                yy = int(yy)
                if yy < 100:
                    yy += 2000
                try:
                    as_of_date = datetime.date(yy, int(mm), int(dd))
                    as_of_label = text
                    found_latest_heading = True
                except ValueError:
                    pass
            continue

        if not found_latest_heading:
            continue  # skip anything before the first dated heading (e.g. "Purchase Order Review")

        text, url = notion_utils.block_plain_text(block)
        if not url or not text:
            continue
        job_match = _JOB_NUMBER_IN_TEXT.search(text)
        if not job_match:
            continue

        jobs.append({
            "job_number": job_match.group(0).upper(),
            "filename": text,
            "url": url,
        })

    return {
        "as_of": as_of_date.isoformat() if as_of_date else None,
        "as_of_label": as_of_label,
        "jobs": jobs,
    }


@app.route("/api/notion/latest-job-analyses", methods=["GET"])
@require_role("admin", "calvin")
def latest_job_analyses():
    try:
        data = _parse_latest_job_analyses()
        return jsonify(data)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Purchase Orders grouped by Job # and a category derived from the actual
# Description text (its first meaningful word) rather than a fixed guess
# list — so grouping reflects whatever terminology is really being used.
# ==========================================

_PO_STOPWORDS = {
    "a", "an", "the", "for", "and", "or", "of", "to", "new", "replacement",
    "part", "parts", "misc", "miscellaneous", "item", "items", "with",
}


def _category_from_description(description):
    if not description:
        return "Other"
    words = re.findall(r"[A-Za-z][A-Za-z\-]*", description)
    for w in words:
        if w.lower() not in _PO_STOPWORDS:
            return w.capitalize()
    return "Other"


@app.route("/api/notion/po-by-job", methods=["GET"])
@require_role("admin", "calvin")
def po_by_job():
    try:
        # Only show POs for jobs that are actually current (per the same
        # "latest heading" definition used for Current Jobs) — otherwise
        # this fills up with completed jobs from long ago.
        latest = _parse_latest_job_analyses()
        current_job_keys = {j["job_number"] for j in latest.get("jobs", [])}

        pages = notion_utils.query_data_source_all(NOTION_PURCHASE_ORDERS_DATASOURCE_ID)
        grouped = {}

        for page in pages:
            props = page.get("properties", {})
            job_number_raw = notion_utils.prop_number(props, "Job #")
            if not job_number_raw:
                continue
            job_key = f"J{int(job_number_raw)}"

            if current_job_keys and job_key not in current_job_keys:
                continue  # not one of the currently active jobs — skip

            name = notion_utils.prop_text(props, "Name") or ""
            description = notion_utils.prop_text(props, "Description") or ""
            category = _category_from_description(description)

            grouped.setdefault(job_key, {}).setdefault(category, []).append({
                "po_number": notion_utils.prop_text(props, "PO #"),
                "vendor": notion_utils.prop_select(props, "Vendor"),
                "name": name,
                "description": description,
                "amount": notion_utils.prop_number(props, "Amount"),
                "status": notion_utils.prop_status(props, "Status"),
                "ordered": notion_utils.prop_date(props, "Ordered"),
                "delivered": notion_utils.prop_date(props, "Delivered"),
                "notion_url": page.get("url"),
            })

        return jsonify({"count": len(grouped), "purchase_orders_by_job": grouped, "current_jobs_considered": sorted(current_job_keys)})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Sales pipeline (Typeform Leads. FEED HERE) — filtered to the statuses
# Jacque said Calvin cares about: NEW, Initial Takeoff, Proposal to review.
# This also covers Kevin's reports, since Kevin is one of the Sales Reps
# feeding this same database.
# ==========================================

_SALES_PIPELINE_STATUSES = ["NEW", "Initial Takeoff", "Proposal to review"]


@app.route("/api/notion/sales-pipeline", methods=["GET"])
@require_role("admin", "calvin")
def sales_pipeline():
    try:
        filter_obj = {
            "and": [
                {"property": "Inactive", "checkbox": {"equals": False}},
                {"or": [
                    {"property": "Project Status", "select": {"equals": s}}
                    for s in _SALES_PIPELINE_STATUSES
                ]},
            ]
        }
        pages = notion_utils.query_data_source(NOTION_SALES_PIPELINE_DATASOURCE_ID, page_size=100, filter_obj=filter_obj)

        items = []
        for page in pages:
            props = page.get("properties", {})
            items.append({
                "project": notion_utils.prop_text(props, "Project"),
                "customer_name": notion_utils.prop_text(props, "Customer Name"),
                "job_number": notion_utils.prop_text(props, "Job #"),
                "project_status": notion_utils.prop_select(props, "Project Status"),
                "sales_rep": notion_utils.prop_select(props, "Sales Rep"),
                "priority": notion_utils.prop_select(props, "Priority"),
                "next_action": notion_utils.prop_text(props, "Next Action"),
                "est_contract": notion_utils.prop_number(props, "Est Contract"),
                "scope_of_work": notion_utils.prop_text(props, "Scope of Work"),
                "notion_url": page.get("url"),
            })

        return jsonify({"count": len(items), "items": items})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Kevin's Daily Management Form — same 7-project-slot shape as the Garage
# "Daily Management Log", but a separate Notion database entirely, synced
# directly from Typeform (not through our script/webhook).
# ==========================================

_KEVIN_FORM_SLOT_FIELDS = [
    ("1 What project did you work on today? ", "1 What did you do for this project? (1)", "1 How long did that take? (2) 1", "1 What is the next action for this project? (2) 1"),
    ("2 What project did you work on today? (3)", "2 What did you do for this project? (2)", "2 How long did that take? (2)", "2 What is the next action for this project? (2)"),
    ("3 What project did you work on today? (4)", "3 What did you do for this project? (3)", "3 How long did that take? (4) 1", "3 What is the next action for this project? (4) 1"),
    ("4 What project did you work on today? (5)", "4 What did you do for this project? (4)", "4 How long did that take? (4)", "4 What is the next action for this project? (4)"),
    ("5 What project did you work on today? (6)", "5 What did you do for this project? (5)", "5 How long did that take? (5)", "5 What is the next action for this project? (6)"),
    ("6 What project did you work on today? (7)", "6 What did you do for this project? (6)", "6 How long did that take? (7)", "6 What is the next action for this project? (7)"),
    ("7 What project did you work on today? (8)", "7 What did you do for this project? (7)", "7 How long did that take? (8)", "7 What is the next action for this project? (8)"),
]


@app.route("/api/notion/kevin-management-form", methods=["GET"])
@require_role("admin", "calvin")
def kevin_management_form():
    try:
        pages = notion_utils.query_data_source(NOTION_KEVIN_FORM_DATASOURCE_ID, page_size=30)
        items = []

        for page in pages:
            props = page.get("properties", {})
            date = notion_utils.prop_date(props, "Date")
            name = notion_utils.prop_people(props, "Name") or "Kevin"

            lines = []
            has_next_action = False
            for project_f, did_f, how_long_f, next_f in _KEVIN_FORM_SLOT_FIELDS:
                project = notion_utils.prop_text(props, project_f)
                did = notion_utils.prop_text(props, did_f)
                how_long = notion_utils.prop_text(props, how_long_f)
                next_action = notion_utils.prop_text(props, next_f)

                if not (project or did):
                    continue

                line = f"{project or 'Project'}: {did or ''}"
                if how_long:
                    line += f" ({how_long})"
                if next_action:
                    line += f" — Next: {next_action}"
                    has_next_action = True
                lines.append(line)

            items.append({
                "id": "kevin-mgmt-" + page["id"],
                "notionUrl": page.get("url"),
                "date": date,
                "name": name,
                "summary": "\n".join(lines),
                "has_next_action": has_next_action,
            })

        items.sort(key=lambda i: i["date"] or "", reverse=True)
        return jsonify({"count": len(items), "items": items})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Self-correcting filename tool — re-derives each PDF's real name/date/job
# directly from its own already-written text content, no external CSV
# needed. Adapted from the team's own original fix_no_date_files.py
# script (structured extraction first, then its proven regex fallback).
#
# Processes in small batches with a continuation cursor rather than all
# at once — with hundreds of PDFs across many months, doing everything in
# one request blew past Vercel's 60s function limit and got killed
# (FUNCTION_INVOCATION_TIMEOUT) before it could even respond. This can no
# longer time out regardless of archive size: the frontend calls this
# repeatedly, passing back the cursor each time, until done.
# ==========================================

def _collect_all_pdf_entries(dbx):
    """Flat list of (folder_path, entry) for every PDF across all year
    folders — just metadata listing, not downloading, so this stays fast
    even for a large archive."""
    entries_flat = []
    try:
        res = dbx.files_list_folder(DROPBOX_BASE_FOLDER)
        base_entries = list(res.entries)
        while res.has_more:
            res = dbx.files_list_folder_continue(res.cursor)
            base_entries.extend(res.entries)
    except Exception:
        return entries_flat

    year_folders = sorted(
        e.name for e in base_entries
        if isinstance(e, dropbox.files.FolderMetadata) and e.name.isdigit() and len(e.name) == 4
    )

    for year in year_folders:
        folder_path = f"{DROPBOX_BASE_FOLDER}/{year}"
        try:
            res = dbx.files_list_folder(folder_path)
            year_entries = list(res.entries)
            while res.has_more:
                res = dbx.files_list_folder_continue(res.cursor)
                year_entries.extend(res.entries)
        except Exception:
            continue
        for e in year_entries:
            if isinstance(e, dropbox.files.FileMetadata) and e.name.lower().endswith(".pdf"):
                entries_flat.append((folder_path, e))

    return entries_flat


def _fix_one_pdf_filename(dbx, folder_path, entry):
    """Downloads, re-parses, and renames a single PDF if needed. Returns
    ('renamed', {...}) / ('skipped_no_change', None) /
    ('skipped_no_date', None) / ('error', message)."""
    try:
        _, res_dl = dbx.files_download(entry.path_display)
        pdf_bytes = res_dl.content
    except Exception as e:
        return ("error", f"{entry.name}: download failed ({e})")

    text = extract_text_from_pdf_bytes(pdf_bytes)

    answers = form_parsers.extract_starred_fields(text)
    has_real_titles = any(a["field"]["title"] for a in answers)

    new_name = None
    new_date = None
    new_job_wo = None

    if answers and has_real_titles:
        name_ans = form_parsers.find_answer(answers, ["your name", "name"])
        date_ans = form_parsers.find_answer(answers, ["date"])
        new_name = form_parsers._clean_or_none(form_parsers.answer_text(name_ans))
        raw_date_text = form_parsers.answer_text(date_ans) or ""
        new_date = format_date_for_jacque(raw_date_text) if raw_date_text else None
        new_job_wo, _pdf_text, _extra = form_parsers.parse_structured_answers(
            answers, new_name or "Unknown", new_date or "No_Date"
        )

    if not new_name or not new_date or new_date == "No_Date":
        raw_date, fallback_name, fallback_job_wo = form_parsers.legacy_regex_extract(text, entry.name)
        if not new_name:
            new_name = fallback_name
        if not new_date or new_date == "No_Date":
            new_date = format_date_for_jacque(raw_date) if raw_date else None
        if not new_job_wo:
            new_job_wo = fallback_job_wo

    if not new_date or new_date == "No_Date" or not new_name:
        return ("skipped_no_date", None)

    parts = [new_date, new_name]
    if new_job_wo:
        parts.append(new_job_wo)
    new_filename = " | ".join(parts) + ".pdf"

    if new_filename == entry.name:
        return ("skipped_no_change", None)

    to_path = f"{folder_path}/{new_filename}"
    try:
        dbx.files_move_v2(entry.path_display, to_path, autorename=True)
        return ("renamed", {"from": entry.name, "to": new_filename})
    except Exception as e:
        return ("error", f"{entry.name}: rename failed ({e})")


@app.route("/api/fix-filenames-from-content", methods=["POST"])
@require_role("admin")
def fix_filenames_from_content():
    try:
        body = request.get_json(silent=True) or {}
        start_index = max(0, int(body.get("start_index", 0)))
        batch_size = 5  # cut way down from 15 — real Dropbox latency in
        # production is apparently higher than local testing could show

        dbx = get_dropbox_client()
        all_entries = _collect_all_pdf_entries(dbx)
        total = len(all_entries)
        batch = all_entries[start_index:start_index + batch_size]

        renamed = []
        skipped_no_change = 0
        skipped_no_date = 0
        errors = []

        # A single malformed/unusual PDF can make pypdf hang indefinitely
        # while parsing it — batching alone only limits how many files run
        # per request, it never protected against ONE file stalling the
        # whole batch forever. Give each file its own hard timeout so a
        # bad file can never block the others or blow the function budget.
        # Note: shutdown(wait=False) is deliberate — a plain 'with' block
        # would still block on exit waiting for every thread to actually
        # finish, even ones already given up on below, defeating the point.
        executor = ThreadPoolExecutor(max_workers=5)
        try:
            future_to_item = {
                executor.submit(_fix_one_pdf_filename, dbx, folder_path, entry): entry
                for folder_path, entry in batch
            }
            for future in future_to_item:
                entry = future_to_item[future]
                try:
                    results_item = future.result(timeout=8)
                except Exception as e:
                    results_item = ("error", f"{entry.name}: timed out or failed ({e})")
                if results_item[0] == "renamed":
                    renamed.append(results_item[1])
                elif results_item[0] == "skipped_no_change":
                    skipped_no_change += 1
                elif results_item[0] == "skipped_no_date":
                    skipped_no_date += 1
                elif results_item[0] == "error":
                    errors.append(results_item[1])
        finally:
            executor.shutdown(wait=False)

        next_index = start_index + len(batch)
        done = next_index >= total or len(batch) == 0

        return jsonify({
            "status": "ok",
            "done": done,
            "next_index": None if done else next_index,
            "total": total,
            "processed_so_far": next_index,
            "renamed_count": len(renamed),
            "renamed": renamed,
            "skipped_no_change": skipped_no_change,
            "skipped_no_date": skipped_no_date,
            "errors": errors,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# Chipperfield Calendar + PTO banner
# ==========================================
# PTO comes from two Notion databases: the current "PTO Request" form (Garage →
# PTO Request Form, Apr 2026 on) and the older "PTO Requests" table (Sep 2025 –
# Apr 2026). Every submitted request is shown (there's no approval field yet),
# including all three types: PTO, Time Off, Sick Time.
NOTION_PTO_DATASOURCE_ID = os.environ.get(
    "NOTION_PTO_DATASOURCE_ID", "3484562c-e56f-808c-a09a-000b4a287b61"
)
NOTION_PTO_OLD_DATASOURCE_ID = os.environ.get(
    "NOTION_PTO_OLD_DATASOURCE_ID", "2654562c-e56f-80b8-8e85-000b24528cea"
)
# 👷 Employee Files (Garage) — active employees' birthdays (month/day only is shown).
NOTION_EMPLOYEE_FILES_DATASOURCE_ID = os.environ.get(
    "NOTION_EMPLOYEE_FILES_DATASOURCE_ID", "47d8f070-51d2-43db-92e4-c016c71796ff"
)
LOG_HISTORY_FIRST_YEAR = 2024
SITE_ID = re.compile(r'^(J|WO)\d{3,}$', re.IGNORECASE)


def _today_central():
    """'Today' for the banner is McCook's date, not the server's (UTC)."""
    try:
        from zoneinfo import ZoneInfo
        return datetime.datetime.now(ZoneInfo("America/Chicago")).date()
    except Exception:
        return (datetime.datetime.utcnow() - datetime.timedelta(hours=6)).date()


def _is_business_day(d):
    return d.weekday() < 5 and not _holiday_name(d)


def _business_day_before(d):
    d = d - datetime.timedelta(days=1)
    while not _is_business_day(d):
        d -= datetime.timedelta(days=1)
    return d


def _multi_select(props, name):
    p = props.get(name) or {}
    if p.get("type") != "multi_select":
        return []
    return [o.get("name") for o in p.get("multi_select", []) if o.get("name")]


def _date_range(props, name):
    p = props.get(name) or {}
    if p.get("type") != "date" or not p.get("date"):
        return None, None
    start = (p["date"].get("start") or "")[:10] or None
    end = (p["date"].get("end") or "")[:10] or start
    return start, end


def _pto_requests():
    """All PTO / Time Off / Sick Time requests from both databases, de-duplicated
    (the old and new tables overlap for a few days in April 2026)."""
    out, seen = [], set()

    def add(name, types, start, end, hours, url):
        if not name or not start:
            return
        key = (name.strip().lower(), start)
        if key in seen:
            return
        seen.add(key)
        if end and end < start:
            start, end = end, start
        out.append({"name": name.strip(), "types": types or ["Time off"], "start": start,
                    "end": end or start, "hours": hours, "url": url})

    for page in notion_utils.query_data_source_all(NOTION_PTO_DATASOURCE_ID):
        props = page.get("properties", {})
        start, end = _date_range(props, "Dates Requested")
        add(notion_utils.prop_text(props, "Employee Name"), _multi_select(props, "What kind of request is this?"),
            start, end, notion_utils.prop_text(props, "Number of hours requested"), page.get("url"))
    try:
        for page in notion_utils.query_data_source_all(NOTION_PTO_OLD_DATASOURCE_ID):
            props = page.get("properties", {})
            start, end = _date_range(props, "Date")
            hours = notion_utils.prop_number(props, "Hours")
            add(notion_utils.prop_text(props, "Employee Name"), _multi_select(props, "Type"),
                start, end, (f"{hours:g}" if hours else None), page.get("url"))
    except Exception:
        pass  # older table is history only; the current form still loads
    return out


def _pto_label(types):
    """'PTO', 'Sick Time', 'Time Off' — or 'PTO / Sick Time' when several were ticked."""
    return " / ".join(types) if types else "Time off"


@app.route("/api/pto-banner", methods=["GET"])
@require_role("admin", "calvin")
def pto_banner():
    """Time off to flag above the daily job logs: shown on the business day
    before it starts, and on its first day."""
    try:
        today = _today_central()
        items = []
        for r in _pto_requests():
            start = datetime.date.fromisoformat(r["start"])
            if today == start or today == _business_day_before(start):
                items.append({**r, "label": _pto_label(r["types"]), "starts_today": today == start})
        items.sort(key=lambda r: (r["start"], r["name"]))
        return jsonify({"status": "ok", "today": today.isoformat(), "items": items})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


def _job_names():
    """Job/WO number → name, from Notion Current Jobs (and Work Orders when reachable)."""
    names = {}
    try:
        for page in notion_utils.query_data_source_all(NOTION_CURRENT_JOBS_DATASOURCE_ID):
            props = page.get("properties", {})
            name = notion_utils.prop_text(props, "Name")
            # "1420" -> "J1420" with the same 4-digit / 5-digit rule the log titles use
            for tag in djl_titles.id_tags(notion_utils.prop_text(props, "Job #") or ""):
                if name:
                    names.setdefault(tag, name)
    except Exception:
        pass
    try:
        for page in notion_utils.query_data_source_all(NOTION_WORKORDERS_DATASOURCE_ID, max_pages=5):
            props = page.get("properties", {})
            wo = notion_utils.prop_number(props, "WO #")
            if not wo:
                continue
            title = next((notion_utils.prop_text(props, k) for k, v in props.items() if v.get("type") == "title"), None)
            label = title or notion_utils.prop_text(props, "Reporting Contact") or ""
            if label:
                names.setdefault(f"WO{int(wo)}", label)
    except Exception:
        pass  # Work Orders source is optional — numbers still show without names
    return names


@app.route("/api/calendar", methods=["GET"])
@require_role("admin", "calvin")
def chipperfield_calendar():
    """
    One month for the Chipperfield Calendar tab. ?month=YYYY-MM (default: this month).
    Days run Monday before the 1st through Sunday after the last day, so the grid is full.
    Per day: job/WO numbers logged that day (from the Daily Job Log file names) with
    the crew, PTO, shop holiday, and birthdays (month/day of active employees).
    Per job: first and last day logged and how many days it's been logged — all history.
    """
    try:
        month = (request.args.get("month") or "").strip()
        today = _today_central()
        try:
            first = datetime.date.fromisoformat(month + "-01") if month else today.replace(day=1)
        except ValueError:
            return jsonify({"status": "error", "message": "month must look like 2026-10"}), 400
        nxt = (first.replace(day=28) + datetime.timedelta(days=4)).replace(day=1)
        grid_start = first - datetime.timedelta(days=first.weekday())
        last = nxt - datetime.timedelta(days=1)
        grid_end = last + datetime.timedelta(days=6 - last.weekday())

        days = {}
        cur = grid_start
        while cur <= grid_end:
            days[cur.isoformat()] = {"sites": {}, "pto": [], "holiday": _holiday_name(cur), "birthdays": []}
            cur += datetime.timedelta(days=1)

        # Job sites, from every log file name (all years), so per-job totals are complete.
        jobs = {}
        for log in list_log_index(range(LOG_HISTORY_FIRST_YEAR, today.year + 1)):
            for site in [s.strip() for s in (log["job_or_wo"] or "").split(",")]:
                if not SITE_ID.match(site):
                    continue
                site = site.upper()
                j = jobs.setdefault(site, {"first": log["date"], "last": log["date"], "dates": set()})
                j["first"] = min(j["first"], log["date"])
                j["last"] = max(j["last"], log["date"])
                j["dates"].add(log["date"])
                if log["date"] in days:
                    entry = days[log["date"]]["sites"].setdefault(site, {"id": site, "crew": [], "logs": []})
                    if log["name"] not in entry["crew"]:
                        entry["crew"].append(log["name"])
                    entry["logs"].append(log["path"])

        # PTO (every day of each request that falls in the grid)
        for r in _pto_requests():
            s = max(datetime.date.fromisoformat(r["start"]), grid_start)
            e = min(datetime.date.fromisoformat(r["end"]), grid_end)
            while s <= e:
                days[s.isoformat()]["pto"].append({"name": r["name"], "label": _pto_label(r["types"]),
                                                   "hours": r["hours"], "start": r["start"], "end": r["end"]})
                s += datetime.timedelta(days=1)

        # Birthdays: month/day only, active employees
        birthdays_ok = True
        try:
            pages = notion_utils.query_data_source_all(NOTION_EMPLOYEE_FILES_DATASOURCE_ID,
                                                       filter_obj={"property": "Status", "select": {"equals": "Active"}})
            for page in pages:
                props = page.get("properties", {})
                bday, _ = _date_range(props, "Birthday")
                name = notion_utils.prop_text(props, "Employee Name")
                if not bday or not name:
                    continue
                m, d = int(bday[5:7]), int(bday[8:10])
                for year in {grid_start.year, grid_end.year}:
                    try:
                        when = datetime.date(year, m, d)
                    except ValueError:  # Feb 29 in a non-leap year
                        when = datetime.date(year, 2, 28)
                    if when.isoformat() in days:
                        days[when.isoformat()]["birthdays"].append(name.split()[0] if name.split() else name)
        except Exception:
            birthdays_ok = False

        names = _job_names()
        visible = {s for d in days.values() for s in d["sites"]}
        job_info = {s: {"name": names.get(s), "first_logged": jobs[s]["first"], "last_logged": jobs[s]["last"],
                        "days_logged": len(jobs[s]["dates"])} for s in visible}
        out_days = [{"date": k, "sites": sorted(v["sites"].values(), key=lambda x: x["id"]),
                     "pto": sorted(v["pto"], key=lambda x: x["name"]), "holiday": v["holiday"],
                     "birthdays": sorted(v["birthdays"])} for k, v in sorted(days.items())]
        return jsonify({"status": "ok", "month": first.strftime("%Y-%m"), "today": today.isoformat(),
                        "days": out_days, "jobs": job_info, "birthdays_loaded": birthdays_ok})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500
